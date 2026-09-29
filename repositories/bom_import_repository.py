"""通用 BOM Excel 源文件适配器。

本模块只负责读取工作簿和按映射识别字段，不负责业务校验或发布事实。
它不会依赖固定文件名、固定工作表名或固定产品。
"""

from __future__ import annotations

import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import xlrd
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_3588M_MAPPING = PROJECT_ROOT / "import_mappings" / "3588m_costing_v1.json"


def _normalized_label(value: object) -> str:
    return re.sub(r"[\s　]+", "", str(value or "")).strip()


def _text(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text or None


@dataclass(frozen=True)
class BomImportMapping:
    mapping_name: str
    metadata_aliases: dict[str, tuple[str, ...]]
    column_aliases: dict[str, tuple[str, ...]]
    required_columns: tuple[str, ...]
    pending_when_absent: tuple[str, ...]
    stop_row_markers: tuple[str, ...]
    total_price_tolerance: str


@dataclass(frozen=True)
class WorkbookSheet:
    name: str
    rows: tuple[tuple[object, ...], ...]


@dataclass(frozen=True)
class MappedBomRow:
    sheet_name: str
    source_row_number: int
    parent_material_code: str | None
    version: str | None
    values: dict[str, object]


@dataclass(frozen=True)
class WorkbookMappingResult:
    rows: tuple[MappedBomRow, ...]
    detected_sheets: tuple[str, ...]
    sheets_without_metadata: tuple[str, ...]


def load_bom_import_mapping(file_path: str | Path = DEFAULT_3588M_MAPPING) -> BomImportMapping:
    mapping_path = Path(file_path)
    payload = json.loads(mapping_path.read_text(encoding="utf-8"))
    return BomImportMapping(
        mapping_name=str(payload["mapping_name"]),
        metadata_aliases={
            name: tuple(str(alias) for alias in definition["aliases"])
            for name, definition in payload["metadata"].items()
        },
        column_aliases={
            name: tuple(str(alias) for alias in definition["aliases"])
            for name, definition in payload["columns"].items()
        },
        required_columns=tuple(str(value) for value in payload["required_columns"]),
        pending_when_absent=tuple(str(value) for value in payload["pending_when_absent"]),
        stop_row_markers=tuple(str(value) for value in payload.get("stop_row_markers", ())),
        total_price_tolerance=str(payload["validation"]["total_price_tolerance"]),
    )


class BomWorkbookRepository:
    """读取并映射旧版/新版 Excel 工作簿。"""

    def read_sheets(self, file_path: str | Path) -> tuple[WorkbookSheet, ...]:
        source_path = Path(file_path)
        suffix = source_path.suffix.lower()
        if suffix == ".xls":
            return self._read_xls(source_path)
        if suffix == ".xlsx":
            return self._read_xlsx(source_path)
        raise ValueError("BOM 导入仅支持 .xls 或 .xlsx 文件")

    def _read_xls(self, file_path: Path) -> tuple[WorkbookSheet, ...]:
        workbook = xlrd.open_workbook(file_path, on_demand=True)
        try:
            return tuple(
                WorkbookSheet(
                    name=sheet_name,
                    rows=tuple(
                        tuple(workbook.sheet_by_name(sheet_name).row_values(row_index))
                        for row_index in range(workbook.sheet_by_name(sheet_name).nrows)
                    ),
                )
                for sheet_name in workbook.sheet_names()
            )
        finally:
            workbook.release_resources()

    def _read_xlsx(self, file_path: Path) -> tuple[WorkbookSheet, ...]:
        workbook = load_workbook(file_path, read_only=True, data_only=True)
        try:
            return tuple(
                WorkbookSheet(
                    name=sheet.title,
                    rows=tuple(tuple(row) for row in sheet.iter_rows(values_only=True)),
                )
                for sheet in workbook.worksheets
            )
        finally:
            workbook.close()

    def map_workbook(
        self,
        file_path: str | Path,
        mapping: BomImportMapping,
    ) -> WorkbookMappingResult:
        mapped_rows: list[MappedBomRow] = []
        detected_sheets: list[str] = []
        sheets_without_metadata: list[str] = []

        for sheet in self.read_sheets(file_path):
            header = self._find_header(sheet, mapping)
            if header is None:
                continue
            header_index, column_indexes = header
            detected_sheets.append(sheet.name)
            metadata = self._extract_metadata(sheet, mapping, header_index)
            if any(not metadata.get(field_name) for field_name in mapping.metadata_aliases):
                sheets_without_metadata.append(sheet.name)

            for zero_based_row_index, row in enumerate(sheet.rows[header_index + 1 :], header_index + 1):
                values = {
                    field_name: (row[column_index] if column_index < len(row) else None)
                    for field_name, column_index in column_indexes.items()
                }
                if not any(value is not None and str(value).strip() for value in values.values()):
                    continue
                normalized_stop_markers = {
                    _normalized_label(marker) for marker in mapping.stop_row_markers
                }
                if any(
                    _normalized_label(value) in normalized_stop_markers
                    for value in values.values()
                    if _text(value) is not None
                ):
                    break
                mapped_rows.append(
                    MappedBomRow(
                        sheet_name=sheet.name,
                        source_row_number=zero_based_row_index + 1,
                        parent_material_code=metadata.get("parent_material_code"),
                        version=metadata.get("version"),
                        values=values,
                    )
                )

        return WorkbookMappingResult(
            rows=tuple(mapped_rows),
            detected_sheets=tuple(detected_sheets),
            sheets_without_metadata=tuple(sheets_without_metadata),
        )

    def _find_header(
        self,
        sheet: WorkbookSheet,
        mapping: BomImportMapping,
    ) -> tuple[int, dict[str, int]] | None:
        alias_lookup = {
            _normalized_label(alias): field_name
            for field_name, aliases in mapping.column_aliases.items()
            for alias in aliases
        }
        for row_index, row in enumerate(sheet.rows):
            column_indexes: dict[str, int] = {}
            for column_index, cell_value in enumerate(row):
                field_name = alias_lookup.get(_normalized_label(cell_value))
                if field_name is not None and field_name not in column_indexes:
                    column_indexes[field_name] = column_index
            if all(field_name in column_indexes for field_name in mapping.required_columns):
                return row_index, column_indexes
        return None

    def _extract_metadata(
        self,
        sheet: WorkbookSheet,
        mapping: BomImportMapping,
        header_index: int,
    ) -> dict[str, str | None]:
        result: dict[str, str | None] = {
            field_name: None for field_name in mapping.metadata_aliases
        }
        rows_to_scan = sheet.rows[: header_index + 1]
        for row in rows_to_scan:
            for cell_index, cell_value in enumerate(row):
                cell_text = _text(cell_value)
                if not cell_text:
                    continue
                for field_name, aliases in mapping.metadata_aliases.items():
                    if result[field_name] is not None:
                        continue
                    for alias in aliases:
                        normalized_cell = _normalized_label(cell_text)
                        normalized_alias = _normalized_label(alias)
                        if normalized_cell == normalized_alias:
                            adjacent = next(
                                (
                                    _text(value)
                                    for value in row[cell_index + 1 :]
                                    if _text(value) is not None
                                ),
                                None,
                            )
                            if adjacent:
                                result[field_name] = adjacent
                                break
                        match = re.search(
                            rf"{re.escape(alias)}\s*[:：]\s*([^\s,，;；]+)",
                            cell_text,
                            flags=re.IGNORECASE,
                        )
                        if match:
                            result[field_name] = match.group(1).strip()
                            break
        return result


__all__ = [
    "BomImportMapping",
    "BomWorkbookRepository",
    "DEFAULT_3588M_MAPPING",
    "MappedBomRow",
    "WorkbookMappingResult",
    "WorkbookSheet",
    "load_bom_import_mapping",
]
