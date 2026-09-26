"""真实库存表的只读 Excel 适配器。"""

from collections.abc import Iterable
from pathlib import Path

from openpyxl import load_workbook

from repositories.normalization import normalize_header, normalize_material_code, normalize_nonnegative_number


RAW_DATA_DIR = Path(__file__).resolve().parent.parent / "data" / "raw"
INVENTORY_FILE = RAW_DATA_DIR / "2026年库存表+料号查询表及成品库存表(1).xlsx"
INVENTORY_SHEET_NAME = "库存表"
RESISTOR_CAPACITOR_SHEET_NAME = "电阻电容现库位"
RESISTOR_CAPACITOR_SHEET_NAME = "电阻电容现库位"
SOURCE_ASSUMPTIONS = {
    "reserved": 0,
    "incoming": 0,
    "safety_stock": "not_provided_in_source",
}


def load_real_inventory_details(file_path: Path = INVENTORY_FILE) -> list[dict[str, object]]:
    """只读取库存表工作表，保留站点、库位和批次库存明细。"""
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        sheet = workbook[INVENTORY_SHEET_NAME]
        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        headers = {normalize_header(value): index for index, value in enumerate(header_row)}
        required_columns = {"归属地", "库位", "物料代码", "批次号", "库存"}
        if not required_columns.issubset(headers):
            raise ValueError("库存表缺少一期必需列")

        details: list[dict[str, object]] = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            material_code = normalize_material_code(row[headers["物料代码"]])
            if not material_code:
                continue
            details.append(
                {
                    "material_code": material_code,
                    "stock": normalize_nonnegative_number(row[headers["库存"]], "库存"),
                    "site": str(row[headers["归属地"]] or "").strip(),
                    "location": str(row[headers["库位"]] or "").strip(),
                    "batch": str(row[headers["批次号"]] or "").strip(),
                    "source_sheet": INVENTORY_SHEET_NAME,
                }
            )
        return details
    finally:
        workbook.close()


def aggregate_real_inventory(details: Iterable[dict[str, object]]) -> dict[str, dict[str, object]]:
    """仅按库存表明细聚合库存；不合并其他库存工作表。"""
    result: dict[str, dict[str, object]] = {}
    for detail in details:
        material_code = str(detail["material_code"])
        stock = normalize_nonnegative_number(detail["stock"], "库存")
        if material_code not in result:
            result[material_code] = {
                "material_code": material_code,
                "stock": 0,
                "reserved": SOURCE_ASSUMPTIONS["reserved"],
                "incoming": SOURCE_ASSUMPTIONS["incoming"],
                "source_assumptions": SOURCE_ASSUMPTIONS.copy(),
                "details": [],
            }
        result[material_code]["stock"] = result[material_code]["stock"] + stock
        result[material_code]["details"].append(detail.copy())
    return result


def load_real_inventory_data(file_path: Path = INVENTORY_FILE) -> dict[str, dict[str, object]]:
    """读取并按物料代码聚合真实库存。"""
    return aggregate_real_inventory(load_real_inventory_details(file_path))


def load_resistor_capacitor_material_codes(file_path: Path = INVENTORY_FILE) -> set[str]:
    """读取待确认来源的料号存在性；不读取其数量进入一期库存。"""
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        sheet = workbook[RESISTOR_CAPACITOR_SHEET_NAME]
        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        headers = {normalize_header(value): index for index, value in enumerate(header_row)}
        if "料号" not in headers:
            raise ValueError("电阻电容现库位缺少料号列")
        return {
            material_code
            for row in sheet.iter_rows(min_row=2, values_only=True)
            if (material_code := normalize_material_code(row[headers["料号"]]))
        }
    finally:
        workbook.close()


def build_real_inventory_scope(
    material_codes: Iterable[str],
    file_path: Path = INVENTORY_FILE,
) -> dict[str, dict[str, object]]:
    """为给定物料生成一期 MRP 所需的库存状态，绝不把未知来源转成零库存。"""
    main_inventory = load_real_inventory_data(file_path)
    excluded_pending_confirmation_codes = load_resistor_capacitor_material_codes(file_path)
    scope: dict[str, dict[str, object]] = {}
    for material_code in material_codes:
        if material_code in main_inventory:
            item = main_inventory[material_code].copy()
            item["inventory_status"] = "known"
            scope[material_code] = item
        elif material_code in excluded_pending_confirmation_codes:
            scope[material_code] = {
                "material_code": material_code,
                "inventory_status": "excluded_pending_confirmation",
                "stock": None,
                "available": None,
                "shortage": None,
                "source_assumptions": {
                    "excluded_source": RESISTOR_CAPACITOR_SHEET_NAME,
                    "reason": "一期库存规则未确认该来源可与主库存合并",
                },
            }
        else:
            scope[material_code] = {
                "material_code": material_code,
                "inventory_status": "missing",
                "stock": None,
                "available": None,
                "shortage": None,
                "source_assumptions": {
                    "reason": "主库存及待确认库存来源均未发现料号",
                },
            }
    return scope


def load_resistor_capacitor_material_codes(file_path: Path = INVENTORY_FILE) -> set[str]:
    """读取待确认来源的料号存在性；不读取其数量进入一期库存。"""
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        sheet = workbook[RESISTOR_CAPACITOR_SHEET_NAME]
        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        headers = {normalize_header(value): index for index, value in enumerate(header_row)}
        if "料号" not in headers:
            raise ValueError("电阻电容现库位缺少料号列")
        return {
            material_code
            for row in sheet.iter_rows(min_row=2, values_only=True)
            if (material_code := normalize_material_code(row[headers["料号"]]))
        }
    finally:
        workbook.close()


def build_real_inventory_scope(
    material_codes: Iterable[str],
    file_path: Path = INVENTORY_FILE,
) -> dict[str, dict[str, object]]:
    """为给定物料生成一期 MRP 所需的库存状态，绝不把未知来源转成零库存。"""
    main_inventory = load_real_inventory_data(file_path)
    excluded_pending_confirmation_codes = load_resistor_capacitor_material_codes(file_path)
    scope: dict[str, dict[str, object]] = {}
    for material_code in material_codes:
        if material_code in main_inventory:
            item = main_inventory[material_code].copy()
            item["inventory_status"] = "known"
            scope[material_code] = item
        elif material_code in excluded_pending_confirmation_codes:
            scope[material_code] = {
                "material_code": material_code,
                "inventory_status": "excluded_pending_confirmation",
                "stock": None,
                "available": None,
                "shortage": None,
                "source_assumptions": {
                    "excluded_source": RESISTOR_CAPACITOR_SHEET_NAME,
                    "reason": "一期库存规则未确认该来源可与主库存合并",
                },
            }
        else:
            scope[material_code] = {
                "material_code": material_code,
                "inventory_status": "missing",
                "stock": None,
                "available": None,
                "shortage": None,
                "source_assumptions": {
                    "reason": "主库存及待确认库存来源均未发现料号",
                },
            }
    return scope
