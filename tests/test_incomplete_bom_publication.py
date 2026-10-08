"""Structure facts can publish with null units; quantity decisions still must refuse.

All writes here are synthetic fixture imports into temporary SQLite databases.
Neither private Excel files nor the configured enterprise database are mutated.
"""

from __future__ import annotations

import json
from datetime import UTC, datetime
from unittest.mock import patch

import pytest
from openpyxl import load_workbook

from fact_layer.models import ImportBatchStatus
from repositories.fact_query_repository import FactQueryRepository
from scripts.import_kq_z112_enterprise_facts import MAPPING, import_enterprise
from services.bom_import_service import BomImportService
from services.fact_query_service import FactQueryService
from services.inventory_import_service import InventoryImportService
from services.production_feasibility_query_service import ProductionFeasibilityQueryService
from tests.test_kq_z112_enterprise_import import enterprise_files, repository, count
from tools.published_product_tools import build_published_product_search


def structure_confirmations(codes):
    return {
        "effective_from": "2025-03-18",
        "parent_unit": None,
        "child_units": {code: None for code in codes},
        "snapshot_at": None,
        "confirmed_by": "夹具审核员",
        "reason": "夹具确认日期及结构发布；计量单位、库存时点仍未知",
    }


def publish_structure(repository, enterprise_files):
    bom, inventory, codes, _ = enterprise_files
    return import_enterprise(
        bom, inventory, confirmations=structure_confirmations(codes),
        publish=True, publication_scope="bom", repository=repository,
    )


def fact_query(repository):
    return FactQueryService(FactQueryRepository(repository))


def test_null_unit_bom_can_publish_without_price(repository, enterprise_files):
    result = publish_structure(repository, enterprise_files)
    assert result["publication_status"] == "published_incomplete"
    assert result["published"] is True
    assert count(repository, "bom_headers") == 1
    assert count(repository, "bom_lines") == 66
    assert count(repository, "price_records") == 0


def test_canonical_parent_is_discoverable_without_guessed_name_unit_or_alias(repository, enterprise_files):
    publish_structure(repository, enterprise_files)
    resolver = FactQueryRepository(repository)
    parent = resolver.resolve_item("KQ-Z112")
    assert parent["material_code"] == "KQ-Z112"
    assert parent["name"] is None
    assert parent["unit"] is None
    assert resolver.resolve_item("969T0Z1120101") is None
    assert count(repository, "item_aliases") == 0


def test_all_66_lines_preserve_original_quantities_and_nulls(repository, enterprise_files):
    publish_structure(repository, enterprise_files)
    result = fact_query(repository).query_bom("KQ-Z112", "V01")
    assert result["status"] == "ok"
    assert result["data"]["revision"] == "V01"
    assert result["data"]["parent_item"]["unit"] is None
    lines = result["data"]["lines"]
    assert len(lines) == 66
    assert {line["child_item"]["material_code"] for line in lines} == set(enterprise_files[2])
    assert all(line["qty_per"] == "2" for line in lines)
    assert all(line["unit"] is None and line["price"] is None for line in lines)
    assert all(line["child_item"]["unit"] is None for line in lines)


def test_query_completeness_explains_structure_only_without_cost(repository, enterprise_files):
    publish_structure(repository, enterprise_files)
    result = fact_query(repository).query_bom("KQ-Z112", "V01")
    assert result["completeness"]["bom_lines"] is True
    assert result["completeness"]["unit"] is False
    assert result["completeness"]["parent_unit"] is False
    assert result["completeness"]["price"] is False
    assert any("unit_missing" in warning for warning in result["warnings"])
    assert any("price_missing" in warning for warning in result["warnings"])
    assert any("结构" in warning and "计算" in warning for warning in result["warnings"])


def test_product_discovery_finds_incomplete_published_bom(repository, enterprise_files):
    publish_structure(repository, enterprise_files)
    result = build_published_product_search(repository.store).invoke({"search_term": "KQ-Z112"})
    assert result["status"] == "ok"
    assert result["data"]["products"][0]["product_code"] == "KQ-Z112"
    assert result["data"]["products"][0]["revisions"][0]["revision"] == "V01"
    assert result["completeness"]["production_feasibility"] is False


@pytest.mark.parametrize("effective_date", [None, "2025-03-18"])
def test_existing_production_safety_gate_refuses_unknown_units(repository, enterprise_files, effective_date):
    publish_structure(repository, enterprise_files)
    service = ProductionFeasibilityQueryService(fact_query(repository))
    result = service.query("KQ-Z112", "V01", requested_quantity=100,
        as_of="2026-09-15T04:00:00+00:00", effective_date=effective_date)
    assert result["status"] == "data_unavailable"
    assert any("单位" in warning for warning in result["warnings"])
    # No numerical feasibility or shortage conclusion can be manufactured.
    data = result["data"] or {}
    assert data.get("inventory_sufficient_for_requested") is None
    assert data.get("max_producible_quantity") is None
    assert data.get("theoretical_max_quantity") is None
    assert not result["completeness"].get("units", False)


def test_no_snapshot_time_keeps_actual_on_hand_rows_in_staging(repository, enterprise_files):
    result = publish_structure(repository, enterprise_files)
    inventory_batch = result["import_batches"]["inventory"]
    assert result["batch_states"]["bom"] == "published"
    assert result["batch_states"]["inventory"] == "pending_confirmation"
    assert "snapshot_at" in repository.get_import_batch(inventory_batch).pending_confirmation_fields
    rows = InventoryImportService(repository).inventory_repository.get_staged_rows(inventory_batch)
    assert len(rows) == 45
    assert all(row.on_hand_qty is not None for row in rows)
    assert all(row.available_qty is None and row.reserved_qty is None for row in rows)
    assert count(repository, "inventory_snapshots") == 0
    assert count(repository, "inventory_facts") == 0


def test_22_missing_inventory_materials_never_become_zero(repository, enterprise_files):
    result = publish_structure(repository, enterprise_files)
    assert result["inventory"]["matched_material_count"] == 44
    assert result["inventory"]["missing_material_count"] == 22
    assert result["inventory"]["missing_reason_counts"] == {"excluded_pending_confirmation": 17, "missing": 5}
    for row in result["inventory"]["missing_materials"]:
        assert row["on_hand_qty"] is None
        assert row["available_qty"] is None
        assert row["shortage_qty"] is None
        inventory = fact_query(repository).query_inventory(row["material_code"])
        assert inventory["status"] == "data_unavailable"


def test_confirmed_snapshot_does_not_promote_on_hand_to_available(repository, enterprise_files):
    result = publish_structure(repository, enterprise_files)
    service = InventoryImportService(repository)
    batch_id = result["import_batches"]["inventory"]
    service.confirm_snapshot_at(batch_id, datetime(2026, 9, 15, 4, tzinfo=UTC))
    service.refresh_item_matches(batch_id)
    service.publish(batch_id)
    first = fact_query(repository).query_inventory(enterprise_files[2][0])
    assert first["data"]["on_hand_qty"] == "23"
    assert first["data"]["available_qty"] is None
    assert first["data"]["reserved_qty"] is None
    assert first["completeness"]["available_qty"] is False
    assert count(repository, "inventory_facts") == 45


def test_structure_capabilities_do_not_claim_deterministic_analysis(repository, enterprise_files):
    result = publish_structure(repository, enterprise_files)
    assert result["capabilities"] == {
        "product_discovery": True,
        "bom_structure_query": True,
        "price_analysis": False,
        "inventory_feasibility": False,
        "deterministic_production_analysis": False,
    }


def test_publication_preserves_header_lines_and_item_provenance(repository, enterprise_files):
    result = publish_structure(repository, enterprise_files)
    query = fact_query(repository).query_bom("KQ-Z112", "V01")
    assert query["provenance"]["source_file"] == enterprise_files[0].name
    assert len(query["provenance"]["content_sha256"]) == 64
    assert query["provenance"]["import_batch_id"] == result["import_batches"]["bom"]
    for expected_row, line in enumerate(query["data"]["lines"], 4):
        assert line["provenance"]["sheet_name"] == "BOM"
        assert line["provenance"]["row_number"] == expected_row
        assert line["provenance"]["import_batch_id"] == result["import_batches"]["bom"]
    assert repository.get_item_audit_history("KQ-Z112")[0]["unit"] is None


@pytest.mark.parametrize("missing", ["effective_from", "confirmed_by", "reason"])
def test_structural_publish_still_requires_confirmed_date_and_audit(enterprise_files, missing):
    confirmations = structure_confirmations(enterprise_files[2])
    confirmations[missing] = None
    with patch("scripts.import_kq_z112_enterprise_facts.create_fact_store", side_effect=AssertionError("must not connect")):
        result = import_enterprise(*enterprise_files[:2], confirmations=confirmations,
            publication_scope="bom", publish=True, migrate=True)
    assert result["status"] == "publish_blocked"
    assert result["published"] is False


def test_confirmation_actor_reason_and_date_are_preserved(repository, enterprise_files):
    result = publish_structure(repository, enterprise_files)
    header = repository.get_staged_header(result["import_batches"]["bom"])
    assert header.parent_unit is None
    assert header.unit_policy == "optional_for_structure"
    assert header.confirmations
    assert any(audit.get("fields", {}).get("effective_from") == "2025-03-18" for audit in header.confirmations)
    assert all(audit["confirmed_by"] == "夹具审核员" for audit in header.confirmations)
    assert all(audit["reason"] == structure_confirmations([])["reason"] for audit in header.confirmations)


def test_staged_structure_can_resume_without_duplicate_batches(repository, enterprise_files):
    bom, inventory, codes, _ = enterprise_files
    staged = import_enterprise(bom, inventory, stage=True, publication_scope="bom", repository=repository)
    result = import_enterprise(bom, inventory, publish=True, publication_scope="bom",
        confirmations=structure_confirmations(codes), repository=repository,
        resume_batches=staged["import_batches"])
    assert result["publication_status"] == "published_incomplete"
    assert count(repository, "import_batches") == 2
    assert count(repository, "bom_lines") == 66
    assert count(repository, "inventory_snapshots") == 0


def test_repeat_version_remains_rejected(repository, enterprise_files):
    publish_structure(repository, enterprise_files)
    with pytest.raises(ValueError, match="已发布"):
        publish_structure(repository, enterprise_files)
    assert count(repository, "bom_headers") == 1
    assert count(repository, "bom_lines") == 66


@pytest.mark.parametrize("quantity", [0, -1, "not-a-quantity"])
def test_unknown_units_never_bypass_illegal_quantity_validation(repository, enterprise_files, quantity):
    bom, _, _, _ = enterprise_files
    workbook = load_workbook(bom)
    workbook["BOM"].cell(4, 4, quantity)
    workbook.save(bom)
    workbook.close()
    staged = BomImportService(repository).stage_bom(bom, MAPPING)
    assert staged.status == ImportBatchStatus.INVALID
    assert staged.errors
    with pytest.raises(ValueError):
        BomImportService(repository).publish(staged.import_batch_id)
    assert count(repository, "items") == 0


def test_optional_price_mapping_keeps_strict_units_unless_explicitly_opted_in(repository, enterprise_files, tmp_path):
    payload = json.loads(MAPPING.read_text(encoding="utf-8"))
    payload.pop("unit_policy", None)
    strict_mapping = tmp_path / "optional_price_strict_units.json"
    strict_mapping.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    service = BomImportService(repository)
    result = service.stage_bom(enterprise_files[0], strict_mapping)
    assert set(result.pending_confirmation_fields) == {"unit", "parent_unit", "effective_from"}
    service.confirm_pending_fields(result.import_batch_id, {"effective_from": "2025-03-18"},
        confirmed_by="夹具审核员", reason="仅确认日期")
    assert repository.get_import_batch(result.import_batch_id).pending_confirmation_fields == ("parent_unit", "unit")
    with pytest.raises(ValueError):
        service.publish(result.import_batch_id)
    assert count(repository, "items") == 0


@pytest.mark.parametrize("effective_date", [None, "2025-03-18"])
def test_confirmed_parent_unit_does_not_bypass_unknown_leaf_units(repository, enterprise_files, effective_date):
    bom, inventory, codes, _ = enterprise_files
    confirmations = structure_confirmations(codes)
    confirmations["parent_unit"] = "台"
    import_enterprise(bom, inventory, confirmations=confirmations, publish=True,
        publication_scope="bom", repository=repository)
    result = ProductionFeasibilityQueryService(fact_query(repository)).query(
        "KQ-Z112", "V01", requested_quantity=100, effective_date=effective_date)
    assert result["status"] == "data_unavailable"
    assert any("单位" in warning for warning in result["warnings"])
    assert (result["data"] or {}).get("inventory_sufficient_for_requested") is None
    assert (result["data"] or {}).get("max_producible_quantity") is None
    assert (result["data"] or {}).get("theoretical_max_quantity") is None


def test_source_effective_date_cannot_bypass_manual_publication_confirmation(repository, enterprise_files):
    bom = enterprise_files[0]
    workbook = load_workbook(bom)
    workbook["BOM"].cell(3, 5, "生效日期")
    for row_number in range(4, 70):
        workbook["BOM"].cell(row_number, 5, "2025-03-18")
    workbook.save(bom)
    workbook.close()
    service = BomImportService(repository)
    result = service.stage_bom(bom, MAPPING)
    assert result.pending_confirmation_fields == ("publication_confirmation",)
    with pytest.raises(ValueError):
        service.publish(result.import_batch_id)
    assert count(repository, "bom_headers") == 0


@pytest.mark.parametrize("missing", ["confirmed_by", "reason"])
def test_confirmation_api_requires_nonempty_audit_even_for_optional_units(repository, enterprise_files, missing):
    service = BomImportService(repository)
    result = service.stage_bom(enterprise_files[0], MAPPING)
    audit = {"confirmed_by": "夹具审核员", "reason": "夹具日期确认"}
    audit[missing] = ""
    with pytest.raises(ValueError):
        service.confirm_pending_fields(result.import_batch_id,
            {"effective_from": "2025-03-18"}, **audit)
    assert repository.get_staged_header(result.import_batch_id).confirmations == ()
    assert count(repository, "bom_headers") == 0


def test_staging_date_tampering_cannot_reuse_another_dates_confirmation(repository, enterprise_files):
    service = BomImportService(repository)
    result = service.stage_bom(enterprise_files[0], MAPPING)
    confirmed = service.confirm_pending_fields(result.import_batch_id,
        {"effective_from": "2025-03-18"}, confirmed_by="夹具审核员", reason="明确确认 2025-03-18")
    assert confirmed.status == ImportBatchStatus.READY_TO_PUBLISH
    # Adversarial corruption is confined to this temporary fixture database.
    with repository.store.transaction() as connection:
        repository.store.execute(connection,
            "UPDATE staged_bom_rows SET effective_from = ? WHERE import_batch_id = ?",
            ("2025-03-19", result.import_batch_id))
    with pytest.raises(ValueError):
        service.publish(result.import_batch_id)
    assert count(repository, "items") == 0
    assert count(repository, "bom_headers") == 0


def test_explicit_alias_resolves_null_unit_without_inventing_one(repository, enterprise_files):
    publish_structure(repository, enterprise_files)
    InventoryImportService(repository).inventory_repository.add_item_alias("KQ-Z112", "EXPLICIT-KQ-ALIAS")
    resolved = FactQueryRepository(repository).resolve_item("EXPLICIT-KQ-ALIAS")
    assert resolved["material_code"] == "KQ-Z112"
    assert resolved["unit"] is None
    query = fact_query(repository).query_bom("EXPLICIT-KQ-ALIAS", "V01")
    assert query["status"] == "ok"
    assert query["data"]["parent_item"]["material_code"] == "KQ-Z112"
    assert query["data"]["parent_item"]["unit"] is None
    assert query["completeness"]["parent_unit"] is False
