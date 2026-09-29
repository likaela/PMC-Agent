"""v1.9.1 事实层存储、审计、流式读取与并发加固测试。"""

from __future__ import annotations

import os
import tempfile
import threading
import unittest
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from uuid import uuid4

from openpyxl import Workbook

from fact_layer.models import ImportBatchStatus, SourceKind
from repositories.bom_import_repository import BomWorkbookRepository, load_bom_import_mapping
from repositories.fact_repository import FactRepository
from repositories.fact_store import PostgreSQLFactStore, SQLiteFactStore, create_fact_store
from repositories.source_file_storage import ImmutableSourceFileStorage, hash_source_file
from scripts.validate_3588m_fact_import import DEFAULT_SOURCE, validate_3588m
from services.bom_import_service import BomImportService


HEADERS = ["序号", "料号", "名称", "规格型号", "数量", "含税单价", "总价"]
POSTGRES_TEST_URL = os.getenv("PMC_TEST_POSTGRESQL_URL", "").strip()


def _create_workbook(
    file_path: Path,
    *,
    parent: str = "3588M",
    version: str = "V-A",
    material_name: str = "原始物料名称",
    row_count: int = 1,
) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "自由命名明细"
    sheet.append(["项目代号", parent, "版本号", version])
    sheet.append([])
    sheet.append([])
    sheet.append(HEADERS)
    for index in range(row_count):
        sheet.append(
            [
                index + 1,
                f"M-{index + 1:04d}",
                material_name if index == 0 else f"物料{index + 1}",
                f"规格{index + 1}",
                index % 3 + 1,
                "1.25",
                str((index % 3 + 1) * 1.25),
            ]
        )
    workbook.save(file_path)
    workbook.close()
    return file_path


def _confirm_and_publish(service: BomImportService, batch_id: str) -> dict[str, object]:
    confirmed = service.confirm_pending_fields(
        batch_id,
        {"unit": "件", "currency": "CNY", "effective_from": "2026-09-29"},
    )
    if confirmed.status != ImportBatchStatus.READY_TO_PUBLISH:
        raise AssertionError("测试批次未进入 ready_to_publish")
    return service.publish(batch_id)


class FactLayerHardeningTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.temp_path = Path(self.temporary_directory.name)
        self.repository = FactRepository(SQLiteFactStore(self.temp_path / "facts.db"))
        self.repository.migrate()
        self.service = BomImportService(self.repository)

    def tearDown(self) -> None:
        self.temporary_directory.cleanup()

    def test_sqlite_memory_database_is_explicitly_rejected(self) -> None:
        with self.assertRaisesRegex(ValueError, "多连接事实存储"):
            SQLiteFactStore(":memory:")
        with self.assertRaisesRegex(ValueError, "多连接事实存储"):
            create_fact_store("sqlite:///:memory:")

    def test_uploaded_source_is_saved_as_immutable_content_addressed_copy(self) -> None:
        source = _create_workbook(self.temp_path / "upload.xlsx")
        storage = ImmutableSourceFileStorage(self.temp_path / "private-sources")

        staged = self.service.stage_bom(
            source,
            source_kind=SourceKind.UPLOAD,
            immutable_storage=storage,
        )
        staged_row = self.repository.get_staged_rows(staged.import_batch_id)[0]
        source_record = self.repository.get_source_file(staged_row.source_file_id)

        self.assertIsNotNone(source_record)
        assert source_record is not None
        self.assertEqual(source_record.source_kind, SourceKind.UPLOAD)
        self.assertTrue(source_record.immutable_copy)
        self.assertIsNone(source_record.source_uri)
        self.assertIsNotNone(source_record.storage_key)
        stored_path = storage.root / str(source_record.storage_key)
        self.assertTrue(stored_path.is_file())
        self.assertEqual(hash_source_file(stored_path)[0], source_record.content_sha256)

    def test_published_source_query_contains_identifier_hash_sheet_row_and_batch(self) -> None:
        source = _create_workbook(self.temp_path / "traceable.xlsx")
        staged = self.service.stage_bom(source)
        _confirm_and_publish(self.service, staged.import_batch_id)

        published = self.service.get_published_bom("3588M", "V-A")

        assert published is not None
        header = published["header"]
        line = published["lines"][0]
        self.assertTrue(header["source_file_id"])
        self.assertEqual(len(header["content_sha256"]), 64)
        self.assertEqual(header["source_kind"], SourceKind.LOCAL.value)
        self.assertTrue(str(header["source_uri"]).startswith(("project://", "file:")))
        self.assertEqual(line["source_file_id"], header["source_file_id"])
        self.assertEqual(line["content_sha256"], header["content_sha256"])
        self.assertEqual(line["source_sheet"], "自由命名明细")
        self.assertEqual(line["source_row_number"], 5)
        self.assertEqual(line["import_batch_id"], staged.import_batch_id)

    def test_large_workbook_only_scans_candidate_boundary_and_matching_bom_rows(self) -> None:
        file_path = self.temp_path / "large-workbook.xlsx"
        workbook = Workbook()
        unrelated = workbook.active
        unrelated.title = "大体量无关数据"
        for index in range(5000):
            unrelated.append([f"无关数据{index}", index])
        candidate = workbook.create_sheet("自由命名BOM")
        candidate.append(["项目代号", "3588M", "版本号", "V-LARGE"])
        candidate.append([])
        candidate.append([])
        candidate.append(HEADERS)
        for index in range(91):
            candidate.append(
                [index + 1, f"L-{index:04d}", "物料", "规格", 1, "1.00", "1.00"]
            )
        candidate.append([None, None, None, "合计", None, None, "91.00"])
        workbook.save(file_path)
        workbook.close()

        mapped = BomWorkbookRepository().map_workbook(file_path, load_bom_import_mapping())

        self.assertEqual(mapped.detected_sheets, ("自由命名BOM",))
        self.assertEqual(len(mapped.rows), 91)
        self.assertLessEqual(mapped.scanned_row_count, 196)

    def test_new_item_revision_does_not_overwrite_old_source_audit(self) -> None:
        first_file = _create_workbook(
            self.temp_path / "first.xlsx",
            parent="PARENT-A",
            material_name="旧名称",
        )
        first = self.service.stage_bom(first_file)
        _confirm_and_publish(self.service, first.import_batch_id)

        second_file = _create_workbook(
            self.temp_path / "second.xlsx",
            parent="PARENT-B",
            material_name="新名称",
        )
        second = self.service.stage_bom(second_file)
        _confirm_and_publish(self.service, second.import_batch_id)

        first_bom = self.service.get_published_bom("PARENT-A", "V-A")
        second_bom = self.service.get_published_bom("PARENT-B", "V-A")
        history = self.repository.get_item_audit_history("M-0001")

        assert first_bom is not None and second_bom is not None
        self.assertEqual(first_bom["lines"][0]["material_name"], "旧名称")
        self.assertEqual(second_bom["lines"][0]["material_name"], "新名称")
        self.assertEqual([record["name"] for record in history], ["旧名称", "新名称"])
        self.assertEqual(
            {record["source_file"] for record in history},
            {"first.xlsx", "second.xlsx"},
        )

    def test_concurrent_duplicate_parent_and_version_has_one_publish_and_one_reject(self) -> None:
        first = self.service.stage_bom(_create_workbook(self.temp_path / "concurrent-a.xlsx"))
        second = self.service.stage_bom(_create_workbook(self.temp_path / "concurrent-b.xlsx"))
        for batch in (first, second):
            self.service.confirm_pending_fields(
                batch.import_batch_id,
                {"unit": "件", "currency": "CNY", "effective_from": "2026-09-29"},
            )

        barrier = threading.Barrier(2)

        def publish(batch_id: str) -> tuple[str, str]:
            barrier.wait()
            try:
                self.service.publish(batch_id)
                return "published", batch_id
            except ValueError as error:
                return "rejected", str(error)

        with ThreadPoolExecutor(max_workers=2) as executor:
            outcomes = list(
                executor.map(publish, (first.import_batch_id, second.import_batch_id))
            )

        self.assertEqual(sorted(outcome[0] for outcome in outcomes), ["published", "rejected"])
        rejected_message = next(value for status, value in outcomes if status == "rejected")
        self.assertIn("已发布", rejected_message)
        self.assertEqual(self.repository.count_published_bom_headers(), 1)

    def test_published_batch_cannot_be_reopened_by_confirmation_update(self) -> None:
        staged = self.service.stage_bom(_create_workbook(self.temp_path / "published.xlsx"))
        _confirm_and_publish(self.service, staged.import_batch_id)

        with self.assertRaisesRegex(ValueError, "状态不允许"):
            self.service.confirm_pending_fields(
                staged.import_batch_id,
                {"unit": "套"},
            )

        batch = self.repository.get_import_batch(staged.import_batch_id)
        assert batch is not None
        self.assertEqual(batch.status, ImportBatchStatus.PUBLISHED)


class LocalPrivate3588MAcceptanceTests(unittest.TestCase):
    @unittest.skipUnless(
        DEFAULT_SOURCE.is_file(),
        "data/raw 私有 3588M 文件不存在，跳过本地验收",
    )
    def test_real_3588m_can_stage_91_traceable_rows(self) -> None:
        result = validate_3588m()

        self.assertEqual(result["result"], "passed")
        self.assertEqual(result["staged_row_count"], 91)
        self.assertEqual(result["error_count"], 0)
        self.assertTrue(result["all_rows_traceable"])


@unittest.skipUnless(
    POSTGRES_TEST_URL,
    "PMC_TEST_POSTGRESQL_URL 未设置，跳过真实 PostgreSQL migration/transaction 集成测试",
)
class PostgreSQLFactStoreIntegrationTests(unittest.TestCase):
    """仅连接调用方明确提供的测试数据库，不使用 mock。"""

    def setUp(self) -> None:
        import psycopg

        self.schema = f"pmc_fact_test_{uuid4().hex}"
        self.admin_connection = psycopg.connect(POSTGRES_TEST_URL, autocommit=True)
        self.admin_connection.execute(f'CREATE SCHEMA "{self.schema}"')
        self.store = PostgreSQLFactStore(POSTGRES_TEST_URL, schema=self.schema)
        self.repository = FactRepository(self.store)

    def tearDown(self) -> None:
        self.admin_connection.execute(f'DROP SCHEMA IF EXISTS "{self.schema}" CASCADE')
        self.admin_connection.close()

    def test_real_postgresql_migrations_publish_and_rollback(self) -> None:
        applied = self.repository.migrate()
        self.assertEqual(
            applied,
            ["001_controlled_fact_layer", "002_source_audit_and_item_history"],
        )

        rollback_id = str(uuid4())
        with self.assertRaises(RuntimeError):
            with self.store.transaction() as connection:
                self.store.execute(
                    connection,
                    "INSERT INTO import_batches "
                    "(id, mapping_name, status, created_at, pending_confirmation_fields_json) "
                    "VALUES (?, ?, ?, ?, ?)",
                    (rollback_id, "rollback", "staging", "2026-09-29T00:00:00+00:00", "[]"),
                )
                raise RuntimeError("force rollback")
        with self.store.transaction() as connection:
            cursor = self.store.execute(
                connection,
                "SELECT COUNT(*) FROM import_batches WHERE id = ?",
                (rollback_id,),
            )
            self.assertEqual(cursor.fetchone()[0], 0)

        with tempfile.TemporaryDirectory() as directory:
            source = _create_workbook(Path(directory) / "postgres.xlsx")
            service = BomImportService(self.repository)
            staged = service.stage_bom(source)
            published = _confirm_and_publish(service, staged.import_batch_id)
            self.assertEqual(published["line_count"], 1)


if __name__ == "__main__":
    unittest.main()
