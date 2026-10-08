"""Generic enterprise structure imports; public tests use only synthetic fixtures."""

import json
import os
import shutil
from pathlib import Path
from unittest.mock import patch

import pytest
from openpyxl import Workbook

from agents.context_resolver import ContextResolver
from agents.planner import AgentPlan, TaskEntities
from repositories.agent_context_repository import AgentContextRepository
from repositories.bom_import_repository import load_bom_import_mapping
from repositories.fact_query_repository import FactQueryRepository
from repositories.fact_repository import FactRepository
from repositories.fact_store import MIGRATIONS_ROOT, SQLiteFactStore
from scripts.import_enterprise_boms import MAPPING, audit_directory, backup_sqlite, stage_audited, publish_staged, published_catalog
from services.bom_import_service import BomImportService
from services.fact_query_service import FactQueryService
from services.production_feasibility_query_service import ProductionFeasibilityQueryService


def workbook(path, product="ENTERPRISE-A", revision="V01", rows=2, priced=False, duplicate=False):
    book = Workbook()
    sheet = book.active
    sheet.title = "自由企业清单"
    sheet.append([f"项目代号:{product} 版本号:{revision}"])
    sheet.append(["料号", "名称", "数量", "含税单价", "总价"])
    for index in range(rows):
        sheet.append(["CHILD-1" if duplicate else f"CHILD-{index+1}", "明确源名称", "2.50",
                      "1.20" if priced else None, "3.00" if priced else None])
    sheet.append([None, "合计", None, None, "6.00"])
    auxiliary = book.create_sheet("辅助说明")
    auxiliary.append(["料号", "数量"])
    auxiliary.append(["DO-NOT-PUBLISH", 999])
    book.save(path)
    book.close()
    return path


@pytest.fixture
def setup(tmp_path):
    raw = tmp_path / "raw"
    raw.mkdir()
    workbook(raw / "name-not-product.xlsx")
    workbook(raw / "second.xlsx", "ENTERPRISE-B", "V02", rows=3, priced=True)
    repository = FactRepository(SQLiteFactStore(tmp_path / "enterprise.db"))
    repository.migrate()
    return raw, repository, tmp_path / "private-report"


def stage(setup):
    raw, repository, output = setup
    audit = audit_directory(raw)
    manifest = stage_audited(audit, raw, repository, output, confirmed_by="test human authorization",
                            reason="explicit structure-only release; no missing field confirmation")
    return audit, manifest


def publish(setup):
    audit, manifest = stage(setup)
    result = publish_staged(manifest, setup[0], setup[1])
    return audit, manifest, result


def test_multiple_products_versions_line_counts_and_original_quantities(setup):
    audit, manifest, result = publish(setup)
    assert audit["counts"] == dict(workbooks=2, bom_files=2, products=2, safe_boms=2)
    assert {(row["parent_material_code"], row["version"], row["line_count"]) for row in result} == {
        ("ENTERPRISE-A", "V01", 2), ("ENTERPRISE-B", "V02", 3)}
    for entry in manifest:
        bom = setup[1].get_published_bom(entry["product"], entry["revision"])
        assert all(str(line["quantity"]) == "2.50" for line in bom["lines"])
        assert bom["header"]["effective_from"] is None
        assert bom["header"]["content_sha256"] == entry["sha256"]
        assert all(line["source_sheet"] == "自由企业清单" and line["source_row_number"] >= 3 for line in bom["lines"])
        assert bom["header"]["import_batch_id"] == entry["import_batch_id"]


def test_staging_does_not_publish_and_keeps_source_price_without_cost_fact(setup):
    _, manifest = stage(setup)
    assert setup[1].count_published_bom_headers() == 0
    priced = next(entry for entry in manifest if entry["product"] == "ENTERPRISE-B")
    rows = setup[1].get_staged_rows(priced["import_batch_id"])
    assert str(rows[0].tax_included_unit_price) == "1.20"
    assert rows[0].currency is None and rows[0].effective_from is None
    with setup[1].store.transaction() as connection:
        assert connection.execute("SELECT COUNT(*) FROM price_records").fetchone()[0] == 0


def test_null_units_prices_dates_completeness_and_capability(setup):
    publish(setup)
    facts = FactQueryService(FactQueryRepository(setup[1]))
    result = facts.query_bom("ENTERPRISE-B", "V02")
    assert result["status"] == "ok"
    assert result["data"]["effective_from"] is None
    assert result["data"]["parent_item"]["unit"] is None
    assert all(line["unit"] is None and line["price"] is None for line in result["data"]["lines"])
    assert result["completeness"]["unit"] is False
    assert result["completeness"]["price"] is False
    assert result["completeness"]["effective_from"] is False
    assert any("effective_date_missing" in warning for warning in result["warnings"])
    assert facts.query_inventory("CHILD-1")["status"] == "data_unavailable"
    assert "on_hand_qty" not in facts.query_inventory("CHILD-1")["data"]
    assert "available_qty" not in facts.query_inventory("CHILD-1")["data"]
    catalog = published_catalog(setup[1])
    assert {item["product"] for item in catalog} == {"ENTERPRISE-A", "ENTERPRISE-B"}
    assert all(not item["capabilities"]["deterministic_production_analysis"] for item in catalog)


@pytest.mark.parametrize("effective_date", [None, "2026-10-08"])
def test_unknown_units_or_period_refuse_production_no_guessed_numbers(setup, effective_date):
    publish(setup)
    facts = FactQueryService(FactQueryRepository(setup[1]))
    result = ProductionFeasibilityQueryService(facts).query("ENTERPRISE-A", "V01", requested_quantity=100, effective_date=effective_date)
    assert result["status"] == "data_unavailable"
    assert result["data"] is None or result["data"].get("max_producible_quantity") is None


def test_structural_bom_query_resolves_unique_version_without_inventory_time(setup):
    publish(setup)
    resolver = ContextResolver(AgentContextRepository(FactQueryRepository(setup[1])))
    plan = AgentPlan(intent="bom", entities=TaskEntities(material_code="ENTERPRISE-A"), required_tools=["bom_query"], missing_parameters=["revision"])
    result = resolver.resolve(plan, "查询ENTERPRISE-A的BOM")
    assert result.status == "ok" and result.plan.entities.revision == "V01"
    assert result.plan.entities.effective_date is None
    assert result.resolutions[0]["provenance"]["selection_scope"] == "structure_only_no_validity_claim"
    assert result.reference()["result"]["provenance"]["bom_candidates"]


def test_explicit_date_confirmation_applies_only_to_selected_product(setup):
    raw, repository, output = setup
    manifest = stage_audited(audit_directory(raw), raw, repository, output, confirmed_by="user",
        reason="confirmed business date", confirmations={"ENTERPRISE-A": {"effective_from": "2025-03-18"}})
    publish_staged(manifest, raw, repository)
    assert repository.get_published_bom("ENTERPRISE-A", "V01")["header"]["effective_from"] == "2025-03-18"
    assert repository.get_published_bom("ENTERPRISE-B", "V02")["header"]["effective_from"] is None


def test_same_product_revision_across_sources_requires_confirmation(setup):
    workbook(setup[0] / "conflict.xlsx", "ENTERPRISE-A")
    audit = audit_directory(setup[0])
    assert all(entry["classification"] == "B" for entry in audit["boms"] if entry["product"] == "ENTERPRISE-A")
    assert audit["counts"]["safe_boms"] == 1


def test_duplicate_children_rejected_not_summed(setup):
    workbook(setup[0] / "duplicate.xlsx", "ENTERPRISE-C", duplicate=True)
    entry = next(row for row in audit_directory(setup[0])["boms"] if row["product"] == "ENTERPRISE-C")
    assert entry["classification"] == "B" and entry["errors"]


@pytest.mark.parametrize("same_cell", [False, True])
def test_conflicting_metadata_not_silently_first_match(setup, same_cell):
    path = setup[0] / "conflicting-metadata.xlsx"
    book = Workbook()
    sheet = book.active
    if same_cell:
        sheet.append(["项目代号:ENTERPRISE-X 项目代号:ENTERPRISE-Y 版本号:V01"])
    else:
        sheet.append(["项目代号:ENTERPRISE-X 版本号:V01"])
        sheet.append(["项目代号:ENTERPRISE-Y 版本号:V01"])
    sheet.append(["料号", "数量"])
    sheet.append(["CHILD-X", 2])
    book.save(path)
    book.close()
    entry = next(row for row in audit_directory(setup[0])["files"] if row["file"] == path.name)
    assert entry["classification"] == "C" and entry["error_type"] == "ValueError"


def test_raw_hash_change_prevents_all_publication(setup):
    _, manifest = stage(setup)
    workbook(setup[0] / manifest[-1]["file"], "CHANGED")
    with pytest.raises(ValueError, match="哈希"):
        publish_staged(manifest, setup[0], setup[1])
    assert setup[1].count_published_bom_headers() == 0


def test_staging_tampering_prevents_all_publication(setup):
    _, manifest = stage(setup)
    with setup[1].store.transaction() as connection:
        connection.execute("UPDATE staged_bom_rows SET quantity='999' WHERE import_batch_id=?", (manifest[-1]["import_batch_id"],))
    with pytest.raises(ValueError, match="明细已变化"):
        publish_staged(manifest, setup[0], setup[1])
    assert setup[1].count_published_bom_headers() == 0


def test_repeat_publication_rejected_and_original_unchanged(setup):
    publish(setup)
    _, manifest = stage(setup)
    with pytest.raises(ValueError, match="已发布"):
        publish_staged(manifest, setup[0], setup[1])
    assert setup[1].count_published_bom_headers() == 2


def test_no_publication_without_explicit_authorization(setup):
    path = setup[0] / "name-not-product.xlsx"
    service = BomImportService(setup[1])
    profile = json.loads(MAPPING.read_text(encoding="utf-8"))
    profile["included_sheets"] = ["自由企业清单"]
    mapping_path = setup[2] / "explicit-sheet.json"
    mapping_path.parent.mkdir(parents=True)
    mapping_path.write_text(json.dumps(profile), encoding="utf-8")
    staged = service.stage_bom(path, mapping_path)
    assert "publication_confirmation" in staged.pending_confirmation_fields
    with pytest.raises(ValueError):
        service.publish(staged.import_batch_id)


def test_backup_is_complete_readonly_and_cannot_overwrite(setup, tmp_path):
    publish(setup)
    target = tmp_path / "backup.db"
    backup = backup_sqlite(Path(setup[1].store.database_path), target)
    assert len(backup["sha256"]) == 64
    assert FactRepository(SQLiteFactStore(target)).count_published_bom_headers() == 2
    with pytest.raises(ValueError):
        backup_sqlite(Path(setup[1].store.database_path), target)


def test_structural_policy_cannot_be_implicit_or_weaken_original_mapping(tmp_path):
    original = load_bom_import_mapping()
    assert original.publication_scope == "full" and original.price_policy == "required" and original.unit_policy == "required"
    profile = json.loads(MAPPING.read_text(encoding="utf-8"))
    profile["price_policy"] = "required"
    path = tmp_path / "bad.json"
    path.write_text(json.dumps(profile), encoding="utf-8")
    with pytest.raises(ValueError):
        load_bom_import_mapping(path)


def test_010_upgrade_preserves_dated_facts_and_fk_with_default_full_scope(tmp_path):
    directory = tmp_path / "legacy" / "sqlite"
    directory.mkdir(parents=True)
    for source in (MIGRATIONS_ROOT / "sqlite").glob("*.sql"):
        if source.stem < "010":
            shutil.copyfile(source, directory / source.name)
    store = SQLiteFactStore(tmp_path / "legacy.db")
    with patch("repositories.fact_store.MIGRATIONS_ROOT", directory.parent):
        store.migrate()
    with store.transaction() as connection:
        connection.execute("INSERT INTO import_batches(id,mapping_name,status,created_at) VALUES('b','old','published','2025-01-01')")
        connection.execute("INSERT INTO source_files(id,import_batch_id,original_filename,content_sha256,byte_size,source_uri) VALUES('s','b','source.xls',?,10,'project://source.xls')", ("a"*64,))
        connection.execute("INSERT INTO bom_headers VALUES('h','OLD','V01','2025-03-18','b','s','BOM',NULL)")
        before = connection.execute("SELECT * FROM bom_headers").fetchall()
    assert store.migrate() == ["010_structural_bom_publication"]
    with store.transaction() as connection:
        assert connection.execute("SELECT * FROM bom_headers").fetchall() == before
        assert connection.execute("PRAGMA foreign_key_check").fetchall() == []
        assert connection.execute("PRAGMA foreign_keys").fetchone()[0] == 1
    assert store.migrate() == []


def test_isolated_demo_catalog_does_not_enter_enterprise_store(setup):
    publish(setup)
    from services.demo_scenario_agent_service import build_demo_scenario_agent
    demo = build_demo_scenario_agent(runtime_root=setup[2] / "isolated-demo-runtime")
    demo_items = FactQueryRepository(FactRepository(demo.repository.store))
    assert demo_items.resolve_item("DEMO-PMC-01") is not None
    assert FactQueryRepository(setup[1]).resolve_item("DEMO-PMC-01") is None
    assert demo_items.resolve_item("ENTERPRISE-A") is None
    assert {entry["product"] for entry in published_catalog(setup[1])} == {"ENTERPRISE-A", "ENTERPRISE-B"}


@pytest.mark.parametrize("message", ["有哪些已发布产品？", "有哪些产品？", "查看产品目录"])
def test_agent_product_catalog_uses_published_facts_and_provenance(setup, message):
    publish(setup)
    from fastapi.testclient import TestClient
    from api.app import app
    with patch.dict(os.environ, {"PMC_FACT_DATABASE_URL": f"sqlite:///{setup[1].store.database_path}", "LLM_PROVIDER": "offline-test"}):
        response = TestClient(app).post("/api/v1/agent/query", json={"message": message, "persist_session": True})
    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] == "ok"
    assert {row["product_code"] for row in payload["structured_data"]["data"]["products"]} == {"ENTERPRISE-A", "ENTERPRISE-B"}
    assert payload["provenance"]["bom_sources"]
    assert payload["scenario_id"] is None


@pytest.mark.parametrize("code,revision,lines", [("ENTERPRISE-A", "V01", 2), ("ENTERPRISE-B", "V02", 3)])
def test_agent_structure_query_without_date_inventory_uses_exact_product(setup, code, revision, lines):
    publish(setup)
    from fastapi.testclient import TestClient
    from api.app import app
    with patch.dict(os.environ, {"PMC_FACT_DATABASE_URL": f"sqlite:///{setup[1].store.database_path}", "LLM_PROVIDER": "offline-test"}):
        response = TestClient(app).post("/api/v1/agent/query", json={"message": f"查询{code}的BOM", "persist_session": True})
    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] == "ok" and payload["intent"] == "bom"
    assert payload["structured_data"]["data"]["revision"] == revision
    assert payload["structured_data"]["data"]["line_count"] == lines
    assert "None" not in payload["answer"]
    assert payload["provenance"]["source_file"]
    assert any("unit_missing" in warning for warning in payload["warnings"])


def test_two_published_versions_require_explicit_structure_choice_without_today(setup):
    workbook(setup[0] / "another-version.xlsx", "ENTERPRISE-A", "V03")
    publish(setup)
    resolver = ContextResolver(AgentContextRepository(FactQueryRepository(setup[1])))
    plan = AgentPlan(intent="bom", entities=TaskEntities(material_code="ENTERPRISE-A"), required_tools=["bom_query"], missing_parameters=["revision"])
    result = resolver.resolve(plan, "查询ENTERPRISE-A的BOM")
    assert result.status == "needs_input" and result.plan.entities.revision is None
    assert {row["revision"] for row in result.candidates} == {"V01", "V03"}
    plan.entities.revision = "V03"
    result = resolver.resolve(plan, "查询ENTERPRISE-A V03 BOM")
    assert result.status == "ok" and result.plan.entities.revision == "V03"


@pytest.mark.parametrize("invalid_qty", [0, -1, "not-number", "NaN"])
def test_structure_policy_still_rejects_invalid_quantities(setup, invalid_qty):
    from openpyxl import load_workbook
    path = workbook(setup[0] / "invalid.xlsx", "ENTERPRISE-INVALID")
    book = load_workbook(path)
    book.active["C3"] = invalid_qty
    book.save(path)
    book.close()
    bom = next(entry for entry in audit_directory(setup[0])["boms"] if entry["product"] == "ENTERPRISE-INVALID")
    assert bom["classification"] == "B" and bom["errors"]


def test_structure_policy_still_validates_original_price_total(setup):
    from openpyxl import load_workbook
    path = workbook(setup[0] / "incorrect-cost.xlsx", "ENTERPRISE-BAD-COST", priced=True)
    book = load_workbook(path)
    book.active["E3"] = "999"
    book.save(path)
    book.close()
    bom = next(entry for entry in audit_directory(setup[0])["boms"] if entry["product"] == "ENTERPRISE-BAD-COST")
    assert bom["classification"] == "B" and bom["errors"]


def test_as_of_production_does_not_use_undated_structural_version(setup):
    publish(setup)
    resolver = ContextResolver(AgentContextRepository(FactQueryRepository(setup[1])))
    plan = AgentPlan(intent="production_capacity", entities=TaskEntities(material_code="ENTERPRISE-A",
        revision="V01", effective_date="2026-10-08", quantity=100), required_tools=["production_analysis"], missing_parameters=[])
    result = resolver.resolve(plan, "生产100台ENTERPRISE-A 生效日期2026-10-08")
    assert result.status == "data_unavailable"
    assert "生效日期尚未确认" in result.message
    assert result.provenance["bom_sources"]


def test_date_confirmation_alone_is_not_structural_publication_authorization(setup):
    service = BomImportService(setup[1])
    profile = json.loads(MAPPING.read_text(encoding="utf-8"))
    profile["included_sheets"] = ["自由企业清单"]
    mapping_path = setup[2] / "scope.json"
    mapping_path.parent.mkdir(parents=True)
    mapping_path.write_text(json.dumps(profile), encoding="utf-8")
    staged = service.stage_bom(setup[0] / "name-not-product.xlsx", mapping_path)
    batch = service.confirm_pending_fields(staged.import_batch_id, {"effective_from": "2025-03-18"},
                                          confirmed_by="business-date-confirmer", reason="date only")
    assert "publication_confirmation" in batch.pending_confirmation_fields
    with pytest.raises(ValueError):
        service.publish(staged.import_batch_id)
