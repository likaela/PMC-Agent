"""v1.9.3 已发布事实受控查询 Tool 测试。"""

from __future__ import annotations

import inspect
import os
import tempfile
import unittest
from datetime import UTC, datetime
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

from fact_layer.models import InboundSupplyInputRow
from repositories.fact_query_repository import FactQueryRepository
from repositories.fact_repository import FactRepository
from repositories.fact_store import SQLiteFactStore
from services.bom_import_service import BomImportService
from services.fact_query_service import FactQueryService
from services.inbound_supply_import_service import InboundSupplyImportService
from services.inventory_import_service import InventoryImportService
from tests.test_inventory_fact_import import _create_inventory_workbook, _seed_item
from tools.agent_tools import ALL_AGENT_TOOLS
from tools.fact_query_tools import (
    FACT_QUERY_TOOLS,
    published_bom_query,
    published_inbound_supply_query,
    published_inventory_batch_query,
    published_inventory_query,
)


def _create_bom_workbook(file_path: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "已发布BOM"
    sheet.append(["项目代号", "PARENT-1", "版本号", "V1"])
    sheet.append([])
    sheet.append([])
    sheet.append(
        [
            "料号",
            "名称",
            "规格型号",
            "数量",
            "含税单价",
            "总价",
            "单位",
            "币种",
            "生效日期",
        ]
    )
    sheet.append(
        ["MAT-001", "查询测试物料", "S1", 2, "3.50", "7.00", "件", "CNY", "2026-09-01"]
    )
    workbook.save(file_path)
    workbook.close()
    return file_path


class FactQueryToolTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.temp_path = Path(self.temporary_directory.name)
        self.database_path = self.temp_path / "facts.db"
        self.fact_repository = FactRepository(SQLiteFactStore(self.database_path))
        self.fact_repository.migrate()
        _seed_item(self.fact_repository, "MAT-001")
        self.inventory_service = InventoryImportService(self.fact_repository)
        self.query_service = FactQueryService(FactQueryRepository(self.fact_repository))
        self.database_environment = patch.dict(
            os.environ,
            {"PMC_FACT_DATABASE_URL": f"sqlite:///{self.database_path.as_posix()}"},
        )
        self.database_environment.start()
        self.addCleanup(self.database_environment.stop)

    def tearDown(self) -> None:
        self.temporary_directory.cleanup()

    def _publish_inventory(
        self,
        name: str,
        snapshot_at: datetime,
        rows: list[list[object]],
    ) -> str:
        source = _create_inventory_workbook(self.temp_path / name, rows)
        staged = self.inventory_service.stage_inventory(source, snapshot_at=snapshot_at)
        self.inventory_service.publish(staged.import_batch_id)
        return staged.import_batch_id

    def test_fact_query_tools_have_separate_registry_without_changing_demo_agent(self) -> None:
        self.assertEqual(
            {query_tool.name for query_tool in FACT_QUERY_TOOLS},
            {
                "published_inventory_query",
                "published_inventory_batch_query",
                "published_bom_query",
                "published_inbound_supply_query",
            },
        )
        current_agent_names = {agent_tool.name for agent_tool in ALL_AGENT_TOOLS}
        self.assertTrue(current_agent_names.isdisjoint({tool.name for tool in FACT_QUERY_TOOLS}))

    def test_canonical_item_inventory_query_returns_structured_totals(self) -> None:
        batch_id = self._publish_inventory(
            "canonical.xlsx",
            datetime(2026, 9, 20, tzinfo=UTC),
            [
                ["S1", "L1", "MAT-001", "B1", 10],
                ["S1", "L2", "MAT-001", "B2", 20],
            ],
        )

        result = published_inventory_query.invoke({"material_code": "MAT-001"})

        self.assertEqual(result["status"], "ok")
        self.assertEqual(result["data"]["canonical_item"]["material_code"], "MAT-001")
        self.assertEqual(result["data"]["on_hand_qty"], "30")
        self.assertEqual(result["data"]["batch_count"], 2)
        self.assertEqual(result["data"]["location_count"], 2)
        self.assertEqual(result["provenance"]["import_batch_id"], batch_id)

    def test_alias_inventory_query_resolves_to_canonical_item(self) -> None:
        self.inventory_service.inventory_repository.add_item_alias("MAT-001", "ALIAS-001")
        self._publish_inventory(
            "alias-query.xlsx",
            datetime(2026, 9, 20, tzinfo=UTC),
            [["S1", "L1", "MAT-001", "B1", 10]],
        )

        result = published_inventory_query.invoke({"material_code": "ALIAS-001"})

        self.assertEqual(result["status"], "ok")
        self.assertEqual(result["query"]["resolved_by"], "item_alias")
        self.assertEqual(result["query"]["resolved_material_code"], "MAT-001")

    def test_unknown_item_returns_item_not_found_not_zero(self) -> None:
        result = published_inventory_query.invoke({"material_code": "UNKNOWN"})

        self.assertEqual(result["status"], "item_not_found")
        self.assertIsNone(result["data"])
        self.assertIsNone(result["provenance"])

    def test_available_and_reserved_unknown_are_not_derived_from_on_hand(self) -> None:
        self._publish_inventory(
            "unknown-quantities.xlsx",
            datetime(2026, 9, 20, tzinfo=UTC),
            [["S1", "L1", "MAT-001", "B1", 10]],
        )

        result = published_inventory_query.invoke({"material_code": "MAT-001"})

        self.assertEqual(result["data"]["on_hand_qty"], "10")
        self.assertIsNone(result["data"]["available_qty"])
        self.assertIsNone(result["data"]["reserved_qty"])
        self.assertFalse(result["completeness"]["available_qty"])
        self.assertFalse(result["completeness"]["reserved_qty"])
        self.assertNotIn("0", {warning for warning in result["warnings"]})

    def test_inventory_batch_query_preserves_rows_and_unknown_fields(self) -> None:
        self._publish_inventory(
            "batch-query.xlsx",
            datetime(2026, 9, 20, tzinfo=UTC),
            [
                ["S1", "L1", "MAT-001", "B1", 10],
                ["S1", "L2", "MAT-001", "B2", 20],
            ],
        )

        result = published_inventory_batch_query.invoke({"material_code": "MAT-001"})

        self.assertEqual(result["status"], "ok")
        self.assertEqual(len(result["data"]["batches"]), 2)
        self.assertEqual(
            {batch["batch_no"] for batch in result["data"]["batches"]}, {"B1", "B2"}
        )
        self.assertTrue(all(batch["received_at"] is None for batch in result["data"]["batches"]))
        self.assertTrue(all(batch["expiry_date"] is None for batch in result["data"]["batches"]))
        self.assertTrue(all(batch["quality_status"] is None for batch in result["data"]["batches"]))
        self.assertTrue(any("不能生成 FIFO" in warning for warning in result["warnings"]))

    def test_as_of_query_selects_latest_snapshot_not_after_requested_time(self) -> None:
        first_batch = self._publish_inventory(
            "snapshot-one.xlsx",
            datetime(2026, 9, 20, tzinfo=UTC),
            [["S1", "L1", "MAT-001", "B1", 10]],
        )
        self._publish_inventory(
            "snapshot-two.xlsx",
            datetime(2026, 9, 25, tzinfo=UTC),
            [["S1", "L1", "MAT-001", "B1", 99]],
        )

        result = published_inventory_query.invoke(
            {"material_code": "MAT-001", "as_of": "2026-09-22T00:00:00+00:00"}
        )
        latest = published_inventory_query.invoke({"material_code": "MAT-001"})

        self.assertEqual(result["status"], "ok")
        self.assertEqual(result["data"]["on_hand_qty"], "10")
        self.assertEqual(result["provenance"]["import_batch_id"], first_batch)
        self.assertEqual(latest["data"]["on_hand_qty"], "99")

    def test_as_of_before_first_snapshot_is_data_unavailable(self) -> None:
        self._publish_inventory(
            "future-snapshot.xlsx",
            datetime(2026, 9, 20, tzinfo=UTC),
            [["S1", "L1", "MAT-001", "B1", 10]],
        )

        result = published_inventory_query.invoke(
            {"material_code": "MAT-001", "as_of": "2026-09-01T00:00:00+00:00"}
        )

        self.assertEqual(result["status"], "data_unavailable")
        self.assertIn("指定时点", result["data"]["reason"])
        self.assertIsNone(result["provenance"])

    def test_published_bom_query_returns_lines_price_and_source_provenance(self) -> None:
        _seed_item(self.fact_repository, "PARENT-1")
        bom_service = BomImportService(self.fact_repository)
        staged = bom_service.stage_bom(_create_bom_workbook(self.temp_path / "bom.xlsx"))
        bom_service.publish(staged.import_batch_id)

        result = published_bom_query.invoke(
            {"parent_material_code": "PARENT-1", "revision": "V1"}
        )

        self.assertEqual(result["status"], "ok")
        self.assertEqual(result["data"]["revision"], "V1")
        self.assertEqual(result["data"]["lines"][0]["qty_per"], "2")
        self.assertEqual(
            result["data"]["lines"][0]["price"]["tax_included_unit_price"], "3.50"
        )
        self.assertEqual(result["provenance"]["import_batch_id"], staged.import_batch_id)
        self.assertEqual(result["data"]["lines"][0]["provenance"]["row_number"], 5)
        self.assertEqual(len(result["provenance"]["content_sha256"]), 64)

    def test_inbound_without_published_authoritative_facts_is_data_unavailable(self) -> None:
        result = published_inbound_supply_query.invoke({"material_code": "MAT-001"})

        self.assertEqual(result["status"], "data_unavailable")
        self.assertEqual(
            result["data"]["reason"],
            "当前事实层没有已发布的权威在途采购数据",
        )
        self.assertIsNone(result["provenance"])

    def test_published_inbound_unknown_received_quantity_stays_unknown(self) -> None:
        source = self.temp_path / "po-source.txt"
        source.write_text("controlled fixture", encoding="utf-8")
        service = InboundSupplyImportService(self.fact_repository)
        staged = service.stage_rows(
            source,
            "authoritative_po_fixture",
            [
                InboundSupplyInputRow(
                    "采购订单",
                    2,
                    "MAT-001",
                    10,
                    received_qty=None,
                    supplier_code="SUP-1",
                    order_no="PO-1",
                )
            ],
            observed_at=datetime(2026, 9, 20, tzinfo=UTC),
        )
        service.publish(staged.import_batch_id)

        result = published_inbound_supply_query.invoke({"material_code": "MAT-001"})

        self.assertEqual(result["status"], "ok")
        self.assertEqual(result["data"]["ordered_qty"], "10")
        self.assertIsNone(result["data"]["received_qty"])
        self.assertIsNone(result["data"]["in_transit_qty"])
        self.assertFalse(result["completeness"]["in_transit_qty"])
        self.assertEqual(result["provenance"]["import_batch_id"], staged.import_batch_id)

    def test_query_tool_source_has_no_raw_excel_reader_or_business_calculation(self) -> None:
        import repositories.fact_query_repository as query_repository_module
        import tools.fact_query_tools as query_tools_module

        source = "\n".join(
            [
                inspect.getsource(query_repository_module),
                inspect.getsource(query_tools_module),
            ]
        ).lower()
        self.assertNotIn("load_workbook", source)
        self.assertNotIn("xlrd", source)
        self.assertNotIn("data/raw", source)
        self.assertNotIn("mrp_service", source)
        self.assertNotIn("production_service", source)

        self._publish_inventory(
            "raw-isolation.xlsx",
            datetime(2026, 9, 20, tzinfo=UTC),
            [["S1", "L1", "MAT-001", "B1", 10]],
        )
        with patch(
            "repositories.inventory_import_repository.load_workbook",
            side_effect=AssertionError("查询 Tool 不应读取 Excel"),
        ):
            result = published_inventory_query.invoke({"material_code": "MAT-001"})
        self.assertEqual(result["status"], "ok")

    def test_all_tool_results_use_the_public_top_level_contract(self) -> None:
        results = [
            published_inventory_query.invoke({"material_code": "UNKNOWN"}),
            published_inventory_batch_query.invoke({"material_code": "UNKNOWN"}),
            published_bom_query.invoke(
                {"parent_material_code": "UNKNOWN", "revision": "V1"}
            ),
            published_inbound_supply_query.invoke({"material_code": "UNKNOWN"}),
        ]
        for result in results:
            self.assertEqual(
                set(result),
                {"status", "query", "data", "completeness", "provenance", "warnings"},
            )


if __name__ == "__main__":
    unittest.main()
