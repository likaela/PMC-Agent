"""v1.9 企业 BOM staging / publish 事实层测试。"""

from __future__ import annotations

import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook

from fact_layer.models import ImportBatchStatus
from repositories.bom_import_repository import BomWorkbookRepository
from repositories.data_source import resolve_data_source_mode
from repositories.fact_repository import FactRepository
from repositories.fact_store import PostgreSQLFactStore, SQLiteFactStore, create_fact_store
from services.bom_import_service import BomImportService


HEADERS = ["序号", "料号", "名称", "规格型号", "数量", "含税单价", "总价"]


class BomImportFactLayerTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.temp_path = Path(self.temporary_directory.name)
        self.fact_repository = FactRepository(SQLiteFactStore(self.temp_path / "facts.db"))
        self.fact_repository.migrate()
        self.service = BomImportService(self.fact_repository)

    def tearDown(self) -> None:
        self.temporary_directory.cleanup()

    def _create_workbook(
        self,
        name: str,
        row_count: int = 1,
        *,
        include_project: bool = True,
        include_version: bool = True,
        include_confirmed_fields: bool = False,
    ) -> Path:
        file_path = self.temp_path / name
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "任意核价明细"
        sheet.cell(row=1, column=1, value="受控 BOM 导入测试")
        if include_project:
            sheet.cell(row=2, column=1, value="项目代号")
            sheet.cell(row=2, column=2, value="3588M")
        if include_version:
            sheet.cell(row=2, column=3, value="版本号")
            sheet.cell(row=2, column=4, value="V-A")
        headers = list(HEADERS)
        if include_confirmed_fields:
            headers.extend(["单位", "币种", "生效日期"])
        for column, header in enumerate(headers, 1):
            sheet.cell(row=4, column=column, value=header)
        for index in range(row_count):
            quantity = index % 3 + 1
            unit_price = "1.25"
            values: list[object] = [
                index + 1,
                f"M-{index + 1:03d}",
                f"物料{index + 1}",
                f"规格{index + 1}",
                quantity,
                unit_price,
                str(quantity * 1.25),
            ]
            if include_confirmed_fields:
                values.extend(["件", "CNY", "2026-09-29"])
            for column, value in enumerate(values, 1):
                sheet.cell(row=5 + index, column=column, value=value)
        workbook.save(file_path)
        workbook.close()
        return file_path

    def _set_cell(self, file_path: Path, cell: str, value: object) -> None:
        workbook = load_workbook(file_path)
        workbook.active[cell] = value
        workbook.save(file_path)
        workbook.close()

    def test_new_bom_with_91_lines_enters_staging_and_requires_confirmation(self) -> None:
        file_path = self._create_workbook("3588M 核价测试.xlsx", row_count=91)

        result = self.service.stage_bom(file_path)

        self.assertEqual(result.staged_row_count, 91)
        self.assertEqual(result.error_count, 0)
        self.assertEqual(result.status, ImportBatchStatus.PENDING_CONFIRMATION)
        self.assertEqual(
            set(result.pending_confirmation_fields),
            {"unit", "currency", "effective_from"},
        )
        self.assertEqual(len(self.fact_repository.get_staged_rows(result.import_batch_id)), 91)

    def test_total_footer_after_91_details_is_not_staged_as_a_bom_line(self) -> None:
        file_path = self._create_workbook("3588M-with-total-footer.xlsx", row_count=91)
        workbook = load_workbook(file_path)
        sheet = workbook.active
        footer_row = 96
        sheet.cell(row=footer_row, column=4, value="合计")
        sheet.cell(row=footer_row, column=7, value="9999.99")
        workbook.save(file_path)
        workbook.close()

        result = self.service.stage_bom(file_path)

        self.assertEqual(result.staged_row_count, 91)
        self.assertEqual(result.error_count, 0)
        self.assertEqual(result.status, ImportBatchStatus.PENDING_CONFIRMATION)

    def test_invalid_rows_are_rejected_with_persisted_reasons_and_no_published_pollution(self) -> None:
        cases = {
            "missing-material.xlsx": ("B5", None, "MATERIAL_CODE_REQUIRED"),
            "nonpositive-quantity.xlsx": ("E5", 0, "QUANTITY_MUST_BE_POSITIVE"),
            "invalid-price.xlsx": ("F5", "not-a-price", "UNIT_PRICE_INVALID"),
        }
        for filename, (cell, value, expected_code) in cases.items():
            with self.subTest(expected_code=expected_code):
                file_path = self._create_workbook(filename)
                self._set_cell(file_path, cell, value)
                result = self.service.stage_bom(file_path)

                self.assertEqual(result.status, ImportBatchStatus.INVALID)
                persisted_codes = {
                    error.error_code
                    for error in self.fact_repository.get_validation_errors(result.import_batch_id)
                }
                self.assertIn(expected_code, persisted_codes)
                with self.assertRaises(ValueError):
                    self.service.publish(result.import_batch_id)
                self.assertEqual(self.fact_repository.count_published_bom_headers(), 0)

    def test_duplicate_bom_line_is_rejected_and_reason_is_retained(self) -> None:
        file_path = self._create_workbook("duplicate.xlsx", row_count=2)
        self._set_cell(file_path, "B6", "M-001")
        result = self.service.stage_bom(file_path)

        errors = self.fact_repository.get_validation_errors(result.import_batch_id)
        duplicate = next(error for error in errors if error.error_code == "DUPLICATE_BOM_LINE")
        self.assertEqual(duplicate.sheet_name, "任意核价明细")
        self.assertEqual(duplicate.source_row_number, 6)
        self.assertEqual(duplicate.raw_value, "M-001")
        self.assertEqual(result.status, ImportBatchStatus.INVALID)

    def test_missing_header_metadata_is_rejected_and_retained(self) -> None:
        file_path = self._create_workbook(
            "missing-metadata.xlsx",
            include_project=False,
            include_version=False,
        )
        result = self.service.stage_bom(file_path)

        errors = self.fact_repository.get_validation_errors(result.import_batch_id)
        self.assertEqual(result.status, ImportBatchStatus.INVALID)
        self.assertEqual(
            {error.field_name for error in errors if error.error_code == "METADATA_REQUIRED"},
            {"parent_material_code", "version"},
        )

    def test_parent_and_version_can_be_read_from_one_metadata_cell(self) -> None:
        file_path = self._create_workbook("combined-metadata.xlsx")
        workbook = load_workbook(file_path)
        sheet = workbook.active
        sheet["A2"] = "项目代号：3588M 版本号：V-COMBINED"
        sheet["B2"] = None
        sheet["C2"] = None
        sheet["D2"] = None
        workbook.save(file_path)
        workbook.close()

        result = self.service.stage_bom(file_path)
        staged_row = self.fact_repository.get_staged_rows(result.import_batch_id)[0]

        self.assertEqual(result.status, ImportBatchStatus.PENDING_CONFIRMATION)
        self.assertEqual(staged_row.parent_material_code, "3588M")
        self.assertEqual(staged_row.version, "V-COMBINED")

    def test_total_price_is_only_an_audit_check(self) -> None:
        file_path = self._create_workbook("total-mismatch.xlsx")
        self._set_cell(file_path, "G5", "999.00")
        result = self.service.stage_bom(file_path)

        codes = {error.error_code for error in result.errors}
        self.assertIn("TOTAL_PRICE_MISMATCH", codes)
        self.assertEqual(result.status, ImportBatchStatus.INVALID)

    def test_confirmed_staging_can_publish_and_query_by_parent_and_version(self) -> None:
        file_path = self._create_workbook("publishable.xlsx", row_count=3)
        staged = self.service.stage_bom(file_path)

        confirmed = self.service.confirm_pending_fields(
            staged.import_batch_id,
            {"unit": "件", "currency": "CNY", "effective_from": "2026-09-29"},
        )
        self.assertEqual(confirmed.status, ImportBatchStatus.READY_TO_PUBLISH)
        published = self.service.publish(staged.import_batch_id)
        queried = self.service.get_published_bom("3588M", "V-A")

        self.assertEqual(published["line_count"], 3)
        self.assertIsNotNone(queried)
        assert queried is not None
        self.assertEqual(queried["header"]["parent_material_code"], "3588M")
        self.assertEqual(queried["header"]["version"], "V-A")
        self.assertEqual(len(queried["lines"]), 3)
        self.assertNotIn("source_total_price", queried["lines"][0])

    def test_published_line_keeps_file_sheet_row_and_batch_lineage(self) -> None:
        file_path = self._create_workbook("traceable-source.xlsx")
        staged = self.service.stage_bom(file_path)
        self.service.confirm_pending_fields(
            staged.import_batch_id,
            {"unit": "件", "currency": "CNY", "effective_from": "2026-09-29"},
        )
        self.service.publish(staged.import_batch_id)

        queried = self.service.get_published_bom("3588M", "V-A")
        assert queried is not None
        line = queried["lines"][0]
        self.assertEqual(line["source_file"], "traceable-source.xlsx")
        self.assertEqual(line["source_sheet"], "任意核价明细")
        self.assertEqual(line["source_row_number"], 5)
        self.assertEqual(line["import_batch_id"], staged.import_batch_id)

    def test_legacy_xls_reader_does_not_require_a_fixed_sheet_name(self) -> None:
        fake_sheet = Mock()
        fake_sheet.nrows = 2
        fake_sheet.row_values.side_effect = [["项目代号", "3588M"], ["料号", "数量"]]
        fake_workbook = Mock()
        fake_workbook.sheet_names.return_value = ["任意旧版工作表"]
        fake_workbook.sheet_by_name.return_value = fake_sheet

        with patch("repositories.bom_import_repository.xlrd.open_workbook", return_value=fake_workbook):
            sheets = BomWorkbookRepository().read_sheets(self.temp_path / "任意文件.xls")

        self.assertEqual(sheets[0].name, "任意旧版工作表")
        fake_workbook.release_resources.assert_called_once()

    def test_postgresql_store_is_selected_for_production_url_without_connecting(self) -> None:
        store = create_fact_store("postgresql://pmc:secret@db.example/pmc")
        self.assertIsInstance(store, PostgreSQLFactStore)

    def test_fact_import_does_not_change_existing_demo_data_source_switch(self) -> None:
        with patch.dict(os.environ, {"PMC_DATA_SOURCE": "demo"}):
            self.assertEqual(resolve_data_source_mode(), "demo")


if __name__ == "__main__":
    unittest.main()
