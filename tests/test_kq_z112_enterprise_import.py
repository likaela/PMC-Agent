"""Enterprise import fixtures live only in temporary databases, never enterprise/demo facts."""

from __future__ import annotations

from dataclasses import replace
from datetime import UTC, datetime
import json
from pathlib import Path
from unittest.mock import patch

import pytest
from openpyxl import Workbook

from fact_layer.models import ImportBatchStatus
from repositories.fact_repository import FactRepository
from repositories.fact_store import SQLiteFactStore
from repositories.fact_query_repository import FactQueryRepository
from repositories.inventory_import_repository import InventoryWorkbookRepository
from services.bom_import_service import BomImportService
from services.fact_query_service import FactQueryService
from services.inventory_import_service import InventoryImportService
from scripts.import_kq_z112_enterprise_facts import (
    MAPPING, DEFAULT_BOM, DEFAULT_INVENTORY, precheck, import_enterprise,
    EnterpriseImportError,
)


@pytest.fixture
def enterprise_files(tmp_path):
    codes = [f"M-{index:03d}" for index in range(1, 67)]
    bom_path = tmp_path / "真实结构测试.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "BOM"
    sheet.append(["来源备注"])
    sheet.append(["项目代号:KQ-Z112       版本号:V01"])
    sheet.append(["料号", "名称", "规格型号", "数量"])
    for code in codes:
        sheet.append([code, "测试物料", "测试规格", 2])
    # A second sheet with a BOM-like table must not become an extra BOM line.
    notes = book.create_sheet("版本说明")
    notes.append(["料号", "名称", "规格型号", "数量"])
    notes.append(["DO-NOT-IMPORT", "说明", "说明", 1])
    book.save(bom_path)
    book.close()
    inv_path = tmp_path / "库存结构测试.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "库存表"
    sheet.append(["归属地", "库位", "物料代码", "批次号", "库存", "权威可用量", "权威预留量"])
    for code in codes[:44]:
        sheet.append(["S1", "L1", code, "B1", 20, 10, 10])
    sheet.append(["S1", "L2", codes[0], "B2", 3, 3, 0])
    sheet.append(["S1", "L9", "OUT-OF-SCOPE", "B9", 999, 999, 0])
    other = book.create_sheet("电阻电容现库位")
    other.append(["料号", "实物数量"])
    for code in codes[44:61]:
        other.append([code, 999])
    master = book.create_sheet("料号查询表")
    master.append(["料号"])
    for code in codes:
        master.append([code])
    book.save(inv_path)
    book.close()
    confirmations = {
        "parent_unit": "台", "effective_from": "2026-09-01", "snapshot_at": "2026-09-15T12:00:00+08:00",
        "child_units": {code: "件" for code in codes}, "confirmed_by": "测试确认人", "reason": "测试夹具业务确认",
    }
    return bom_path, inv_path, codes, confirmations


@pytest.fixture
def repository(tmp_path):
    repository = FactRepository(SQLiteFactStore(tmp_path / "test-facts.db"))
    repository.migrate()
    return repository


def publish_fixture(repository, enterprise_files):
    bom, inv, _, confirmations = enterprise_files
    return import_enterprise(bom, inv, confirmations=confirmations, publish=True, repository=repository)


def count(repository, table):
    with repository.store.transaction() as connection:
        return repository.store.execute(connection, f"SELECT COUNT(*) FROM {table}").fetchone()[0]


def test_default_precheck_never_connects_or_migrates(enterprise_files):
    bom, inv, _, _ = enterprise_files
    with patch("scripts.import_kq_z112_enterprise_facts.create_fact_store", side_effect=AssertionError("must not connect")):
        result = import_enterprise(bom, inv)
    assert not result["published"]
    assert result["bom"]["line_count"] == 66
    assert result["pending_confirmation_fields"]
    assert result["readiness_report"]["publication_executed"] is False
    assert result["readiness_report"]["bom"]["effective_date"]["status"] == "pending_confirmation"


def test_structure_staging_requires_date_and_audit_without_guessing_units(repository, enterprise_files):
    service = BomImportService(repository)
    result = service.stage_bom(enterprise_files[0], MAPPING)
    assert result.error_count == 0
    assert result.staged_row_count == 66
    assert result.status == ImportBatchStatus.PENDING_CONFIRMATION
    assert set(result.pending_confirmation_fields) == {"effective_from", "publication_confirmation"}
    assert result.completeness["price"] is False
    assert result.completeness["unit"] is False
    assert result.warnings
    assert all(row.tax_included_unit_price is None for row in repository.get_staged_rows(result.import_batch_id))
    with pytest.raises(ValueError):
        service.publish(result.import_batch_id)
    assert count(repository, "items") == 0


def test_kq_parent_and_version_and_66_lines_are_published(repository, enterprise_files):
    result = publish_fixture(repository, enterprise_files)
    query = FactQueryRepository(repository)
    parent = query.resolve_item("KQ-Z112")
    assert parent["material_code"] == "KQ-Z112"
    assert parent["name"] is None
    assert parent["unit"] == "台"
    bom = repository.get_published_bom("KQ-Z112", "V01")
    assert bom["header"]["version"] == "V01"
    assert len(bom["lines"]) == 66
    assert result["published"]
    assert count(repository, "items") == 67
    assert count(repository, "item_aliases") == 0
    assert query.resolve_item("969T0Z1120101") is None


def test_unpriced_published_query_has_null_price_completeness_and_warning(repository, enterprise_files):
    publish_fixture(repository, enterprise_files)
    result = FactQueryService(FactQueryRepository(repository)).query_bom("KQ-Z112", "V01")
    assert result["status"] == "ok"
    assert len(result["data"]["lines"]) == 66
    assert all(line["price"] is None for line in result["data"]["lines"])
    assert result["completeness"]["price"] is False
    assert result["warnings"]
    assert count(repository, "price_records") == 0
    assert all(line["tax_included_unit_price"] is None for line in repository.get_published_bom("KQ-Z112", "V01")["lines"])


def test_parent_and_lines_keep_source_provenance(repository, enterprise_files):
    publish_fixture(repository, enterprise_files)
    audit = repository.get_item_audit_history("KQ-Z112")[0]
    assert audit["source_row_number"] == 2
    assert audit["source_sheet"] == "BOM"
    assert audit["source_file"] == enterprise_files[0].name
    assert len(audit["content_sha256"]) == 64
    result = FactQueryService(FactQueryRepository(repository)).query_bom("KQ-Z112", "V01")
    assert result["provenance"]["source_file"] == enterprise_files[0].name
    assert [line["provenance"]["row_number"] for line in result["data"]["lines"]] == list(range(4, 70))


def test_44_matches_and_22_missing_with_reasons_no_zero(enterprise_files):
    bom, inv, _, _ = enterprise_files
    result = precheck(bom, inv)
    inventory = result["inventory"]
    assert inventory["matched_material_count"] == 44
    assert inventory["missing_material_count"] == 22
    assert inventory["missing_reason_counts"] == {"excluded_pending_confirmation": 17, "missing": 5}
    assert all(row["on_hand_qty"] is None and row["shortage_qty"] is None for row in inventory["missing_materials"])
    assert result["material_master"]["bom_matched_count"] == 66
    readiness = result["readiness_report"]
    assert readiness["inventory"]["on_hand_coverage"]["known_material_count"] == 44
    assert readiness["inventory"]["on_hand_coverage"]["percentage"] == "66.67%"
    assert readiness["inventory"]["available_coverage"]["known_material_count"] == 0
    assert readiness["inventory"]["available_coverage"]["field_status"] == "not_provided"
    assert not readiness["production_analysis"]["can_generate_deterministic_conclusion"]


def test_inventory_null_does_not_become_zero_or_on_hand(repository, enterprise_files):
    publish_fixture(repository, enterprise_files)
    service = FactQueryService(FactQueryRepository(repository))
    first = service.query_inventory(enterprise_files[2][0])
    assert first["data"]["on_hand_qty"] == "23"
    assert first["data"]["available_qty"] is None
    assert first["data"]["reserved_qty"] is None
    assert not first["completeness"]["available_qty"]
    assert first["data"]["batch_count"] == 2
    missing = service.query_inventory(enterprise_files[2][-1])
    assert missing["status"] == "data_unavailable"
    assert count(repository, "inventory_facts") == 45


def test_explicit_available_columns_are_preserved(repository, enterprise_files):
    bom, inv, codes, confirmations = enterprise_files
    confirmations["inventory_optional_columns"] = {
        "available_qty": ["权威可用量"], "reserved_qty": ["权威预留量"]}
    import_enterprise(bom, inv, confirmations=confirmations, publish=True, repository=repository)
    result = FactQueryService(FactQueryRepository(repository)).query_inventory(codes[0])
    assert result["data"]["on_hand_qty"] == "23"
    assert result["data"]["available_qty"] == "13"
    assert result["data"]["reserved_qty"] == "10"
    assert result["completeness"]["available_qty"]


def test_duplicate_version_rejected_without_additional_facts(repository, enterprise_files):
    publish_fixture(repository, enterprise_files)
    with pytest.raises(ValueError, match="已发布"):
        publish_fixture(repository, enterprise_files)
    assert count(repository, "bom_headers") == 1
    assert count(repository, "bom_lines") == 66
    assert count(repository, "items") == 67


def test_publish_flag_cannot_override_missing_business_confirmations(enterprise_files):
    with patch("scripts.import_kq_z112_enterprise_facts.create_fact_store", side_effect=AssertionError("must not connect")):
        result = import_enterprise(enterprise_files[0], enterprise_files[1], publish=True, migrate=True)
    assert result["status"] == "publish_blocked"
    assert not result["published"]


def test_staged_batches_resume_without_duplicate_staging(repository, enterprise_files):
    bom, inv, _, confirmations = enterprise_files
    staged = import_enterprise(bom, inv, stage=True, repository=repository)
    assert staged["batch_states"] == {"bom": "pending_confirmation", "inventory": "pending_confirmation"}
    result = import_enterprise(bom, inv, confirmations=confirmations, publish=True,
        repository=repository, resume_batches=staged["import_batches"])
    assert result["published"]
    assert count(repository, "import_batches") == 2
    assert count(repository, "staged_bom_rows") == 66
    header = repository.get_staged_header(staged["import_batches"]["bom"])
    assert header.confirmations
    assert all(row["confirmed_by"] == "测试确认人" for row in header.confirmations)


def test_snapshot_timezone_is_required(enterprise_files):
    bom, inv, _, confirmations = enterprise_files
    confirmations["snapshot_at"] = "2026-09-15T12:00:00"
    with pytest.raises(ValueError, match="时区"):
        precheck(bom, inv, confirmations)


def test_declared_illegal_price_is_not_ignored_in_optional_mode(repository, enterprise_files, tmp_path):
    from openpyxl import load_workbook
    book = load_workbook(enterprise_files[0])
    book["BOM"].cell(3, 5, "含税单价")
    book["BOM"].cell(4, 5, "not-a-price")
    book.save(enterprise_files[0])
    book.close()
    result = BomImportService(repository).stage_bom(enterprise_files[0], MAPPING)
    assert result.status == ImportBatchStatus.INVALID
    assert any(error.error_code == "UNIT_PRICE_INVALID" for error in result.errors)
    with pytest.raises(ValueError):
        BomImportService(repository).publish(result.import_batch_id)
    assert count(repository, "items") == 0


@pytest.mark.parametrize("columns", [
    {"available_qty": "库存"}, {"available_qty": ["库存"]},
    {"available_qty": ["权威量"], "reserved_qty": ["权威量"]},
])
def test_optional_mapping_cannot_reinterpret_stock_or_overlap(columns):
    with pytest.raises(ValueError):
        InventoryWorkbookRepository(optional_column_aliases=columns)


def test_precheck_detects_source_changes(enterprise_files):
    from scripts.import_kq_z112_enterprise_facts import _source
    bom, inv, _, _ = enterprise_files
    bom_source, inv_source = _source(bom), _source(inv)
    with patch("scripts.import_kq_z112_enterprise_facts._source", side_effect=[
        bom_source, inv_source, {**bom_source, "sha256": "changed"},
    ]):
        with pytest.raises(ValueError, match="源文件发生变化"):
            precheck(bom, inv)


def test_changed_file_after_staging_cannot_publish_and_batches_remain_auditable(repository, enterprise_files):
    from openpyxl import load_workbook
    bom, inv, _, confirmations = enterprise_files
    original = InventoryImportService.stage_inventory

    def stage_and_change(service, *args, **kwargs):
        result = original(service, *args, **kwargs)
        book = load_workbook(bom)
        book["BOM"].cell(1, 1, "更改后的来源备注")
        book.save(bom)
        book.close()
        return result

    with patch.object(InventoryImportService, "stage_inventory", stage_and_change):
        with pytest.raises(EnterpriseImportError, match="源文件发生变化") as caught:
            import_enterprise(bom, inv, confirmations=confirmations, publish=True, repository=repository)
    assert set(caught.value.report["import_batches"]) == {"bom", "inventory"}
    assert not caught.value.report["published"]
    assert count(repository, "bom_headers") == 0
    assert count(repository, "items") == 0
    assert count(repository, "inventory_facts") == 0


def test_partial_failure_report_can_resume_without_republishing_bom(repository, enterprise_files):
    bom, inv, _, confirmations = enterprise_files
    with patch.object(InventoryImportService, "publish", side_effect=RuntimeError("simulated storage failure")):
        with pytest.raises(EnterpriseImportError) as caught:
            import_enterprise(bom, inv, confirmations=confirmations, publish=True, repository=repository)
    report = caught.value.report
    assert report["status"] == "import_failed"
    assert report["batch_states"]["bom"] == "published"
    assert report["batch_states"]["inventory"] != "published"
    assert "simulated storage failure" not in report["failure_reason"]
    result = import_enterprise(bom, inv, confirmations=confirmations, publish=True,
        repository=repository, resume_batches=report["import_batches"])
    assert result["published"]
    assert result["bom_publication"]["resumed_existing_batch"]
    assert count(repository, "import_batches") == 2
    assert count(repository, "bom_headers") == 1
    assert count(repository, "inventory_facts") == 45


def test_resume_requires_original_enterprise_import_policy(repository, enterprise_files):
    bom, inv, _, confirmations = enterprise_files
    staged = import_enterprise(bom, inv, stage=True, repository=repository)
    header = repository.get_staged_header(staged["import_batches"]["bom"])
    with patch.object(repository, "get_staged_header", return_value=replace(header, publish_parent_item=False)):
        with pytest.raises(EnterpriseImportError, match="身份不匹配"):
            import_enterprise(bom, inv, confirmations=confirmations, publish=True,
                repository=repository, resume_batches=staged["import_batches"])
    assert count(repository, "bom_headers") == 0


def test_optional_price_strict_units_does_not_guess_parent_unit(repository, enterprise_files, tmp_path):
    # Optional pricing never implicitly enables incomplete structural publication.
    mapping = json.loads(MAPPING.read_text(encoding="utf-8"))
    mapping["unit_policy"] = "required"
    strict_mapping = tmp_path / "strict_units_mapping.json"
    strict_mapping.write_text(json.dumps(mapping, ensure_ascii=False), encoding="utf-8")
    service = BomImportService(repository)
    result = service.stage_bom(enterprise_files[0], strict_mapping)
    state = service.confirm_pending_fields(result.import_batch_id, {"unit": "件", "effective_from": "2026-09-01"},
        confirmed_by="测试确认人", reason="测试确认")
    assert state.pending_confirmation_fields == ("parent_unit",)
    with pytest.raises(ValueError):
        service.publish(result.import_batch_id)


def test_migration_is_additive_and_idempotent(repository):
    assert repository.migrate() == []
    with repository.store.transaction() as connection:
        tables = {row[0] for row in repository.store.execute(connection,
            "SELECT name FROM sqlite_master WHERE type='table'").fetchall()}
    assert {"staged_bom_headers", "price_records", "inventory_facts", "agent_sessions"} <= tables


@pytest.mark.skipif(not DEFAULT_BOM.is_file() or not DEFAULT_INVENTORY.is_file(), reason="私有 KQ 原始文件不存在")
def test_private_kq_precheck_is_read_only():
    from repositories.source_file_storage import hash_source_file
    before = [hash_source_file(path) for path in (DEFAULT_BOM, DEFAULT_INVENTORY)]
    result = precheck(DEFAULT_BOM, DEFAULT_INVENTORY)
    assert result["bom"]["line_count"] == 66
    assert result["inventory"]["matched_material_count"] == 44
    assert result["inventory"]["missing_material_count"] == 22
    assert result["inventory"]["missing_reason_counts"] == {"excluded_pending_confirmation": 17, "missing": 5}
    assert not result["errors"]
    assert not result["published"]
    assert result["readiness_report"]["bom"]["line_count"] == 66
    assert result["readiness_report"]["inventory"]["available_coverage"]["known_material_count"] == 0
    assert not result["readiness_report"]["price"]["can_calculate_cost"]
    assert [hash_source_file(path) for path in (DEFAULT_BOM, DEFAULT_INVENTORY)] == before
