"""KQ enterprise import. Default: read-only precheck; writes require --stage/--publish.

No demo fallback, SQL writes, guessed units/dates, prices or available quantities.
Reports contain private material codes and belong in ignored data/normalized/.
"""

from __future__ import annotations

import argparse
import json
import sys
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import load_workbook
from repositories.bom_import_repository import BomWorkbookRepository, load_bom_import_mapping, _text
from repositories.fact_repository import FactRepository
from repositories.fact_store import create_fact_store
from repositories.inventory_import_repository import InventoryWorkbookRepository
from repositories.source_file_storage import hash_source_file
from services.bom_import_service import BomImportService
from services.inventory_import_service import InventoryImportService, _quantity
from services.enterprise_data_readiness_service import build_enterprise_readiness_report

MAPPING = PROJECT_ROOT / "import_mappings/kq_z112_bom_v1.json"
DEFAULT_BOM = PROJECT_ROOT / "data/raw/10 KQ-Z112 BOM 969T0Z1120101  23.xls"
DEFAULT_INVENTORY = PROJECT_ROOT / "data/raw/2026年库存表+料号查询表及成品库存表(1).xlsx"
DEFAULT_REPORT = PROJECT_ROOT / "data/normalized/kq_z112_enterprise_import_report.json"


class EnterpriseImportError(ValueError):
    """Preserve batch identifiers for recovery without claiming cross-batch atomicity."""

    def __init__(self, reason: str, report: dict):
        super().__init__(reason)
        self.report = report


def _snapshot(value: object) -> datetime | None:
    if value is None:
        return None
    parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    if parsed.tzinfo is None or parsed.utcoffset() is None:
        raise ValueError("库存快照时间必须包含明确时区")
    return parsed


def _source(path: Path) -> dict:
    digest, size = hash_source_file(path)
    return {"filename": path.name, "sha256": digest, "byte_size": size}


def _verify_sources(report: dict, bom_path: Path, inventory_path: Path) -> None:
    for key, path in (("bom", bom_path), ("inventory", inventory_path)):
        if _source(path) != report["sources"][key]:
            raise ValueError("导入期间源文件发生变化，请停止并核对来源")


def _sheet_codes(workbook, sheet_name: str, code_header: str) -> set[str]:
    if sheet_name not in workbook.sheetnames:
        return set()
    rows = workbook[sheet_name].iter_rows(max_col=64, values_only=True)
    headers = next(rows, ())
    normalized = [str(value or "").replace(" ", "").strip() for value in headers]
    if code_header not in normalized:
        return set()
    column = normalized.index(code_header)
    return {code for row in rows if (code := _text(row[column])) is not None}


def precheck(bom_path: Path, inventory_path: Path, confirmations: dict | None = None) -> dict:
    """No database, migration, source copies or staged writes are performed here."""
    confirmations = confirmations or {}
    if not isinstance(confirmations, dict):
        raise ValueError("确认文件必须是字段字典")
    sources = {"bom": _source(bom_path), "inventory": _source(inventory_path)}
    allowed = {"effective_from", "parent_unit", "child_units", "snapshot_at", "confirmed_by",
               "reason", "inventory_optional_columns"}
    if set(confirmations) - allowed:
        raise ValueError("确认文件含未支持字段；不能通过确认文件提供价格或伪造库存")
    if not isinstance(confirmations.get("child_units", {}), dict):
        raise ValueError("child_units 必须是逐料号单位字典")
    mapping = load_bom_import_mapping(MAPPING)
    mapped = BomWorkbookRepository().map_workbook(bom_path, mapping)
    identities = {(row.parent_material_code, row.version) for row in mapped.rows}
    if identities != {("KQ-Z112", "V01")} or len(mapped.rows) != 66:
        raise ValueError("源文件元数据及明细必须确认为 KQ-Z112 / V01 / 66 条；不从文件名推断")
    validator = BomImportService(None)
    duplicates = set()
    errors = []
    scope = set()
    validated_rows = []
    for row in mapped.rows:
        staged, row_errors = validator._validate_row(
            "precheck", "precheck", row, mapping, Decimal(mapping.total_price_tolerance), duplicates)
        scope.add(staged.material_code)
        validated_rows.append(staged)
        errors.extend({"sheet": error.sheet_name, "row": error.source_row_number,
                       "code": error.error_code, "field": error.field_name,
                       "message": error.message} for error in row_errors)
    scope.discard(None)
    child_units = confirmations.get("child_units", {})
    if set(child_units) - scope:
        raise ValueError("单位确认包含不属于当前 BOM 的料号")
    for unit in [*child_units.values(), confirmations.get("parent_unit")]:
        if unit is not None and (not isinstance(unit, str) or not unit.strip()):
            raise ValueError("计量单位须为明确的非空字符串")
    effective = confirmations.get("effective_from")
    if effective is not None:
        date.fromisoformat(str(effective))
    snapshot = _snapshot(confirmations.get("snapshot_at"))
    optional_columns = confirmations.get("inventory_optional_columns", {})
    inventory = InventoryWorkbookRepository(optional_column_aliases=optional_columns).map_workbook(inventory_path)
    main_codes = {_text(row.values.get("material_code")) for row in inventory.rows}
    main_codes.discard(None)
    if inventory.detected_sheet is None:
        errors.append({"code": "INVENTORY_SHEET_NOT_FOUND", "message": "未找到已确认的库存表"})
    inventory_rows = [row for row in inventory.rows if _text(row.values.get("material_code")) in scope
                      or _text(row.values.get("material_code")) is None]
    for row in inventory_rows:
        if _text(row.values.get("material_code")) is None:
            errors.append({"sheet": row.sheet_name, "row": row.source_row_number,
                           "code": "MATERIAL_CODE_REQUIRED", "message": "库存行缺少料号，无法判断归属"})
        for field in ("on_hand_qty", "available_qty", "reserved_qty"):
            try:
                value = _quantity(row.values.get(field))
                if (field == "on_hand_qty" and value is None) or (value is not None and value < 0):
                    raise InvalidOperation
            except InvalidOperation:
                errors.append({"sheet": row.sheet_name, "row": row.source_row_number,
                               "code": "INVENTORY_QUANTITY_INVALID", "field": field,
                               "message": "库存数量非法"})
    workbook = load_workbook(inventory_path, read_only=True, data_only=True)
    try:
        other_sources = {name: _sheet_codes(workbook, name, header) for name, header in (
            ("电阻电容现库位", "料号"), ("成品库存表", "物料代码"),
            ("2023年至2025年形成的存货-可以处理", "物料代码"))}
        master = _sheet_codes(workbook, "料号查询表", "料号")
    finally:
        workbook.close()
    missing = []
    for code in sorted(scope - main_codes):
        excluded = [name for name, codes in other_sources.items() if code in codes]
        missing.append({"material_code": code,
            "inventory_status": "excluded_pending_confirmation" if excluded else "missing",
            "excluded_sources": excluded,
            "reason": "仅在排除来源有记录，未完成时点及去重确认" if excluded
                      else "主库存及本次审计的其他库存来源均无记录，不能当作零库存",
            "on_hand_qty": None, "available_qty": None, "shortage_qty": None})
    pending = []
    if effective is None and any(row.effective_from is None for row in validated_rows):
        pending.append("effective_from")
    if confirmations.get("parent_unit") is None:
        pending.append("parent_unit")
    missing_units = sorted(code for code in scope if child_units.get(code) is None)
    # This workbook has no unit/effective date fields; confirmations must cover its rows.
    source_units = {_text(row.values.get("material_code")): _text(row.values.get("unit")) for row in mapped.rows}
    missing_units = [code for code in missing_units if not source_units.get(code)]
    if missing_units:
        pending.append("child_units")
    if any("currency" in row.pending_fields for row in validated_rows):
        pending.append("currency")
    if snapshot is None:
        pending.append("snapshot_at")
    for field in ("confirmed_by", "reason"):
        if not isinstance(confirmations.get(field), str) or not confirmations[field].strip():
            pending.append(field)
    result = {
        "status": "invalid" if errors else "pending_confirmation" if pending else "precheck_passed",
        "published": False, "product": "KQ-Z112", "revision": "V01",
        "sources": sources,
        "bom": {"line_count": len(mapped.rows), "metadata_sheet": mapped.rows[0].sheet_name,
                "metadata_row": mapped.rows[0].metadata_row_number,
                "price_complete": all(_text(row.values.get("tax_included_unit_price")) is not None for row in mapped.rows)},
        "inventory": {"raw_row_count": inventory.source_row_count, "scope_material_count": len(scope),
            "matched_material_count": len(scope & main_codes), "selected_row_count": len(inventory_rows),
            "excluded_out_of_scope_row_count": len(inventory.rows) - len(inventory_rows),
            "is_full_workbook_import": False,
            "missing_material_count": len(missing), "missing_materials": missing,
            "missing_reason_counts": dict(Counter(row["inventory_status"] for row in missing)),
            "available_complete": not missing and bool(inventory_rows)
                 and all(row.values.get("available_qty") is not None for row in inventory_rows),
            "reserved_complete": not missing and bool(inventory_rows)
                 and all(row.values.get("reserved_qty") is not None for row in inventory_rows),
            "snapshot_at": snapshot.isoformat() if snapshot else None,
            "excluded_sources": list(inventory.excluded_sources)},
        "material_master": {"bom_matched_count": len(scope & master), "published_separately": False},
        "pending_confirmation_fields": pending, "missing_unit_material_codes": missing_units,
        "errors": errors,
        "warnings": ["没有价格时不创建 PriceRecord", "在库不等于可用，缺失库存不等于零",
                     "仅导入 BOM 范围主库存明细，不是整份库存工作簿发布",
                     "供应商主数据不用于推断采购参数或供应关系"],
    }
    result["readiness_report"] = build_enterprise_readiness_report(
        product=result["product"], revision=result["revision"],
        bom_lines=[{
            "material_code": row.material_code, "quantity": str(row.quantity) if row.quantity is not None else None,
            "validation_status": row.validation_status,
            "unit": child_units.get(row.material_code) or row.unit,
            "effective_from": str(effective) if effective is not None else
                              row.effective_from.isoformat() if row.effective_from else None,
            "unit_price": str(row.tax_included_unit_price) if row.tax_included_unit_price is not None else None,
            "currency": row.currency,
        } for row in validated_rows],
        parent_unit=confirmations.get("parent_unit"),
        inventory_rows=[{"material_code": _text(row.values.get("material_code")),
            **{field: row.values.get(field) for field in ("on_hand_qty", "available_qty", "reserved_qty")}}
            for row in inventory_rows],
        snapshot_at=snapshot.isoformat() if snapshot else None,
        effective_date_source="business_confirmation" if effective is not None else "source_workbook",
        available_field_present=any("available_qty" in row.values for row in inventory_rows),
        validation_errors=errors, provenance=sources,
    )
    # Structural publication and production readiness are separate contracts.
    # Unit gaps remain visible in pending/completeness, but are not invented to publish.
    bom_required = [field for field in pending if field in {
        "effective_from", "currency", "confirmed_by", "reason"}]
    result["publication_requirements"] = {
        "bom": bom_required,
        "inventory": [field for field in pending if field == "snapshot_at"],
        "all": [*bom_required, *(["snapshot_at"] if snapshot is None else [])],
    }
    result["publication_status"] = "not_published"
    result["bom"]["completeness"] = {
        "parent_unit": confirmations.get("parent_unit") is not None,
        "unit": not missing_units,
        "price": result["bom"]["price_complete"],
    }
    if confirmations.get("parent_unit") is None or missing_units:
        result["warnings"].append("unit_missing：单位未知保持 null，不阻止结构发布，但禁止确定性数量计算")
    if not result["bom"]["price_complete"]:
        result["warnings"].append("price_missing：没有已发布价格，不能计算成本")
    result["capabilities"] = {
        "product_discovery": False, "bom_structure_query": False,
        "price_analysis": False, "inventory_feasibility": False,
        "deterministic_production_analysis": False,
    }
    _verify_sources(result, bom_path, inventory_path)
    return result


def import_enterprise(bom_path: Path, inventory_path: Path, *, confirmations: dict | None = None,
                      stage: bool = False, publish: bool = False, migrate: bool = False,
                      database_url: str | None = None, repository: FactRepository | None = None,
                      resume_batches: dict | None = None, publication_scope: str = "all") -> dict:
    if publication_scope not in {"bom", "all"}:
        raise ValueError("publication_scope 必须为 bom 或 all")
    confirmations = confirmations or {}
    report = precheck(bom_path, inventory_path, confirmations)
    report["publication_scope"] = publication_scope
    if not stage and not publish:
        return report
    if report["errors"] or (publish and report["publication_requirements"][publication_scope]):
        report["status"] = "publish_blocked" if publish else "invalid"
        return report  # Do not create a database/batch when preconditions for publication fail.
    repository = repository or FactRepository(create_fact_store(database_url))
    if migrate:
        report["applied_migrations"] = repository.migrate()
    with repository.store.transaction() as connection:
        exists = repository.store.execute(connection,
            "SELECT version FROM schema_migrations WHERE version = ?", ("009_incomplete_bom_units",)).fetchone()
    if not exists:
        raise ValueError("请显式执行 009_incomplete_bom_units 迁移，或在写入命令中传入 --migrate")
    bom = BomImportService(repository)
    columns = confirmations.get("inventory_optional_columns", {})
    inventory = InventoryImportService(repository,
        InventoryWorkbookRepository(optional_column_aliases=columns))
    ids = dict(resume_batches or {})
    report["import_batches"] = ids
    try:
        _execute_import(repository, bom, inventory, report, ids, bom_path, inventory_path,
                        confirmations, publish=publish, publication_scope=publication_scope)
    except Exception as error:
        reason = str(error) if isinstance(error, (ValueError, KeyError)) else "导入存储操作失败，请核对批次后重试"
        report.update(status="import_failed", published=False, failure_reason=reason)
        report["batch_states"] = {}
        for key, batch_id in ids.items():
            try:
                batch = repository.get_import_batch(batch_id)
                report["batch_states"][key] = batch.status.value if batch else "not_found"
            except Exception:
                report["batch_states"][key] = "state_unavailable"
        raise EnterpriseImportError(reason, report) from error
    return report


def _execute_import(repository, bom, inventory, report, ids, bom_path, inventory_path,
                    confirmations, *, publish, publication_scope):
    """Use the existing import services; each publication remains its own transaction."""
    if ids:
        if set(ids) != {"bom", "inventory"}:
            raise ValueError("继续导入必须提供 BOM 和库存两项批次标识")
        header = repository.get_staged_header(ids["bom"])
        inv_rows = inventory.inventory_repository.get_staged_rows(ids["inventory"])
        if (header is None or not inv_rows or (header.parent_material_code, header.version) != ("KQ-Z112", "V01")
                or header.price_policy != "optional" or header.unit_policy != "optional_for_structure"
                or not header.publish_parent_item):
            raise ValueError("继续导入批次身份不匹配")
        for key, source_id in (("bom", header.source_file_id), ("inventory", inv_rows[0].source_file_id)):
            source = repository.get_source_file(source_id)
            if source is None or source.content_sha256 != report["sources"][key]["sha256"]:
                raise ValueError("源文件已变化，禁止继续旧批次")
        scope_report = inventory.get_quality_report(ids["inventory"]) or {}
        codes = {row.material_code for row in repository.get_staged_rows(ids["bom"])}
        if set(scope_report.get("scope", {}).get("material_codes", [])) != codes:
            raise ValueError("继续导入的库存范围与 BOM 不一致")
    else:
        result = bom.stage_bom(bom_path, MAPPING)
        ids["bom"] = result.import_batch_id
    bom_state = repository.get_import_batch(ids["bom"])
    if bom_state.status.value != "published":
        actor = confirmations.get("confirmed_by")
        reason = confirmations.get("reason")
        if actor and reason:
            if confirmations.get("parent_unit"):
                bom.confirm_parent_unit(ids["bom"], confirmations["parent_unit"], confirmed_by=actor, reason=reason)
            if confirmations.get("effective_from"):
                bom.confirm_pending_fields(ids["bom"], {"effective_from": confirmations["effective_from"]},
                    confirmed_by=actor, reason=reason)
            for code, unit in confirmations.get("child_units", {}).items():
                if unit is None:
                    continue  # Explicitly unknown, not a confirmation and never a default unit.
                bom.confirm_pending_fields(ids["bom"], {"unit": unit}, code, confirmed_by=actor, reason=reason)
    codes = {row.material_code for row in repository.get_staged_rows(ids["bom"])}
    if "inventory" not in ids:
        result = inventory.stage_inventory(inventory_path,
            snapshot_at=_snapshot(confirmations.get("snapshot_at")), material_scope=codes)
        ids["inventory"] = result.import_batch_id
    elif confirmations.get("snapshot_at") and repository.get_import_batch(ids["inventory"]).status.value != "published":
        inventory.confirm_snapshot_at(ids["inventory"], _snapshot(confirmations["snapshot_at"]))
    report["batch_states"] = {key: repository.get_import_batch(value).status.value for key, value in ids.items()}
    _verify_sources(report, bom_path, inventory_path)
    header = repository.get_staged_header(ids["bom"])
    inv_rows = inventory.inventory_repository.get_staged_rows(ids["inventory"])
    if header is None or not inv_rows:
        raise ValueError("BOM 或库存暂存明细为空，禁止发布")
    for key, source_id in (("bom", header.source_file_id), ("inventory", inv_rows[0].source_file_id)):
        source = repository.get_source_file(source_id)
        if source is None or source.content_sha256 != report["sources"][key]["sha256"]:
            raise ValueError("暂存来源与预检不一致，禁止发布")
    if publish:
        if any(repository.get_import_batch(batch_id).error_count for batch_id in ids.values()):
            raise ValueError("暂存批次存在校验错误，禁止发布")
        if repository.get_import_batch(ids["bom"]).status.value != "published":
            report["bom_publication"] = bom.publish(ids["bom"])
        else:
            report["bom_publication"] = {"status": "published", "resumed_existing_batch": True}
        if repository.get_import_batch(ids["inventory"]).status.value != "published":
            # Match against the newly published Items, while keeping facts in staging.
            inventory.refresh_item_matches(ids["inventory"])
        if publication_scope == "all":
            inventory_state = repository.get_import_batch(ids["inventory"])
            if inventory_state.status.value != "published":
                report["inventory_publication"] = inventory.publish(ids["inventory"])
            else:
                report["inventory_publication"] = {"status": "published", "resumed_existing_batch": True}
        else:
            report["inventory_publication"] = {
                "status": "not_published", "reason": "本次仅发布 BOM 结构事实，库存保留原批次状态",
                "snapshot_at": report["inventory"]["snapshot_at"],
            }
        report["published"] = True
    report["inventory_quality_report"] = inventory.get_quality_report(ids["inventory"])
    report["batch_states"] = {key: repository.get_import_batch(value).status.value for key, value in ids.items()}
    report["status"] = "published" if report["published"] else "staged"
    if report["published"]:
        published_bom = repository.get_published_bom("KQ-Z112", "V01")
        report["bom"]["completeness"] = published_bom["completeness"]
        report["publication_status"] = (
            "published_incomplete" if published_bom["warnings"] or
                report["batch_states"]["inventory"] != "published" or
                not report["inventory"]["available_complete"] else "published")
        report["capabilities"].update(product_discovery=True, bom_structure_query=True,
            price_analysis=published_bom["completeness"]["price"] and published_bom["completeness"]["unit"])
        if publication_scope == "bom":
            report["status"] = report["publication_status"]


def _write_report(report_path: Path, result: dict) -> None:
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description="KQ-Z112 企业事实受控导入，默认只读预检")
    parser.add_argument("--bom", type=Path, default=DEFAULT_BOM)
    parser.add_argument("--inventory", type=Path, default=DEFAULT_INVENTORY)
    parser.add_argument("--report", type=Path, default=DEFAULT_REPORT)
    parser.add_argument("--confirmations", type=Path)
    parser.add_argument("--database-url")
    parser.add_argument("--migrate", action="store_true")
    parser.add_argument("--resume-report", type=Path)
    parser.add_argument("--publication-scope", choices=("bom", "all"), default="all",
                        help="bom 仅发布结构事实；all 另外要求确认快照时间并发布库存")
    modes = parser.add_mutually_exclusive_group()
    modes.add_argument("--stage", action="store_true")
    modes.add_argument("--publish", action="store_true")
    args = parser.parse_args()
    raw_root = (PROJECT_ROOT / "data/raw").resolve()
    report_path = args.report.resolve()
    inputs = {args.bom.resolve(), args.inventory.resolve()}
    if args.confirmations:
        inputs.add(args.confirmations.resolve())
    if args.resume_report:
        inputs.add(args.resume_report.resolve())
    if report_path.is_relative_to(raw_root) or report_path in inputs:
        parser.error("报告不得写入 raw 目录或覆盖输入文件")
    if (args.migrate or args.resume_report) and not (args.stage or args.publish):
        parser.error("迁移或继续批次必须配合显式 --stage / --publish")
    from dotenv import load_dotenv
    load_dotenv(PROJECT_ROOT / ".env")
    try:
        confirmations = json.loads(args.confirmations.read_text(encoding="utf-8")) if args.confirmations else {}
        resume = json.loads(args.resume_report.read_text(encoding="utf-8"))["import_batches"] if args.resume_report else None
        result = import_enterprise(args.bom, args.inventory, confirmations=confirmations,
            stage=args.stage, publish=args.publish, migrate=args.migrate,
            database_url=args.database_url, resume_batches=resume, publication_scope=args.publication_scope)
    except EnterpriseImportError as error:
        _write_report(report_path, error.report)
        print(json.dumps({"status": "import_failed", "reason": str(error),
                          "batch_states": error.report["batch_states"]}, ensure_ascii=False))
        return 1
    except (ValueError, KeyError, OSError) as error:
        print(json.dumps({"status": "failed", "reason": str(error)}, ensure_ascii=False))
        return 1
    _write_report(report_path, result)
    print(json.dumps({"status": result["status"], "published": result["published"],
        "bom_line_count": result["bom"]["line_count"],
        "inventory_match": f'{result["inventory"]["matched_material_count"]}/{result["inventory"]["scope_material_count"]}',
        "missing_inventory_count": result["inventory"]["missing_material_count"],
        "publication_status": result["publication_status"], "capabilities": result["capabilities"],
        "pending_confirmation_fields": result["pending_confirmation_fields"]}, ensure_ascii=False))
    return 1 if result["status"] in {"invalid", "publish_blocked"} else 0


if __name__ == "__main__":
    raise SystemExit(main())
