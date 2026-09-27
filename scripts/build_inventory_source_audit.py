"""只读比对真实 BOM、库存来源与料号主数据，生成库存来源口径审计。"""

from __future__ import annotations

import argparse
import json
from collections.abc import Iterable
from pathlib import Path

import xlrd
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parent.parent
RAW_DATA_DIR = PROJECT_ROOT / "data" / "raw"
DOC_OUTPUT = PROJECT_ROOT / "docs" / "inventory_source_audit.md"
BOM_FILE = RAW_DATA_DIR / "10 KQ-Z112 BOM 969T0Z1120101  23.xls"
INVENTORY_FILE = RAW_DATA_DIR / "2026年库存表+料号查询表及成品库存表(1).xlsx"


def normalize_code(value: object) -> str:
    """将 Excel 中的料号标准化为稳定字符串。"""
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


def normalize_header(value: object) -> str:
    return str(value or "").replace(" ", "").replace("　", "").strip()


def numeric_or_zero(value: object) -> int | float:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return int(value) if float(value).is_integer() else float(value)
    return 0


def read_bom() -> list[dict[str, str]]:
    workbook = xlrd.open_workbook(str(BOM_FILE))
    sheet = workbook.sheet_by_name("BOM")
    headers = {normalize_header(value): index for index, value in enumerate(sheet.row_values(2))}
    records: list[dict[str, str]] = []
    for row_index in range(3, sheet.nrows):
        material_code = normalize_code(sheet.cell_value(row_index, headers["料号"]))
        if material_code:
            records.append(
                {
                    "material_code": material_code,
                    "bom_material_name": str(sheet.cell_value(row_index, headers["名称"])).strip(),
                }
            )
    return records


def read_sheet_by_code(
    sheet_name: str,
    code_header: str,
    stock_header: str,
) -> dict[str, dict[str, object]]:
    """读取指定工作表，按料号汇总指定数量字段，并保留出现标记。"""
    workbook = load_workbook(INVENTORY_FILE, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name]
        first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        headers = {normalize_header(value): index for index, value in enumerate(first_row)}
        if code_header not in headers or stock_header not in headers:
            raise ValueError(f"{sheet_name} 缺少 {code_header} 或 {stock_header}")

        result: dict[str, dict[str, object]] = {}
        for row in sheet.iter_rows(min_row=2, values_only=True):
            code = normalize_code(row[headers[code_header]])
            if not code:
                continue
            current = result.setdefault(code, {"stock": 0, "row_count": 0})
            current["stock"] = current["stock"] + numeric_or_zero(row[headers[stock_header]])
            current["row_count"] = current["row_count"] + 1
        return result
    finally:
        workbook.close()


def read_material_master() -> dict[str, str]:
    workbook = load_workbook(INVENTORY_FILE, read_only=True, data_only=True)
    try:
        sheet = workbook["料号查询表"]
        first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        headers = {normalize_header(value): index for index, value in enumerate(first_row)}
        records: dict[str, str] = {}
        for row in sheet.iter_rows(min_row=2, values_only=True):
            code = normalize_code(row[headers["料号"]])
            if code and code not in records:
                records[code] = str(row[headers["名称"]] or "").strip()
        return records
    finally:
        workbook.close()


def sum_available_sources(sources: Iterable[dict[str, dict[str, object]]], code: str) -> int | float | None:
    matches = [source[code]["stock"] for source in sources if code in source]
    if not matches:
        return None
    total = sum(matches)
    return int(total) if float(total).is_integer() else total


def build_audit() -> dict[str, object]:
    bom = read_bom()
    main = read_sheet_by_code("库存表", "物料代码", "库存")
    resistor_capacitor = read_sheet_by_code("电阻电容现库位", "料号", "实物数量")
    finished_goods = read_sheet_by_code("成品库存表", "物料代码", "库存")
    historical = read_sheet_by_code("2023年至2025年形成的存货-可以处理", "物料代码", "库存数量")
    master = read_material_master()
    other_sources = [resistor_capacitor, finished_goods, historical]

    records: list[dict[str, object]] = []
    missing_main: list[dict[str, object]] = []
    for bom_record in bom:
        code = bom_record["material_code"]
        in_main = code in main
        in_resistor = code in resistor_capacitor
        in_finished = code in finished_goods
        in_historical = code in historical
        in_other = in_resistor or in_finished or in_historical
        source_names = [
            name
            for name, present in (
                ("电阻电容现库位", in_resistor),
                ("成品库存表", in_finished),
                ("2023年至2025年形成的存货-可以处理", in_historical),
            )
            if present
        ]
        if in_main:
            mrp_status = "主库存已匹配，参与一期 MRP"
        elif in_other:
            mrp_status = "主库存未匹配，其他来源有记录，未参与一期 MRP"
        else:
            mrp_status = "主库存及其他审计来源均无记录，MRP 库存数据缺失"

        record = {
            "material_code": code,
            "material_name": master.get(code) or bom_record["bom_material_name"],
            "in_main_inventory_sheet": in_main,
            "in_resistor_capacitor_sheet": in_resistor,
            "in_finished_goods_sheet": in_finished,
            "in_historical_inventory": in_historical,
            "in_material_master": code in master,
            "main_inventory_stock": main[code]["stock"] if in_main else None,
            "other_source_stock": sum_available_sources(other_sources, code),
            "current_mrp_status": mrp_status,
        }
        records.append(record)

        if not in_main:
            if in_other:
                duplicate_risk = "是：其他库存来源与主库存表存在 BOM 料号交集，未完成来源归属与时点对账"
                safe_to_include = "否"
                reason = (
                    "仅出现在 "
                    + "、".join(source_names)
                    + "。该来源未与主库存完成归属、时点和去重对账；按一期规则不自动并入可用库存。"
                )
            else:
                duplicate_risk = "否：四类审计来源均未发现该料号"
                safe_to_include = "否"
                reason = "库存表、电阻电容现库位、成品库存表和历史可处理库存均无记录，不能将缺失数据当作库存 0 或可用库存。"
            missing_main.append(
                {
                    "material_code": code,
                    "material_name": record["material_name"],
                    "other_inventory_sources": "、".join(source_names) if source_names else "无",
                    "other_source_stock": record["other_source_stock"],
                    "possible_duplicate_inventory_source": duplicate_risk,
                    "safe_for_phase_one_available_stock": safe_to_include,
                    "reason": reason,
                }
            )

    in_main_count = sum(record["in_main_inventory_sheet"] for record in records)
    in_resistor_count = sum(record["in_resistor_capacitor_sheet"] for record in records)
    in_finished_count = sum(record["in_finished_goods_sheet"] for record in records)
    in_historical_count = sum(record["in_historical_inventory"] for record in records)
    all_source_count = sum(
        record["in_main_inventory_sheet"]
        or record["in_resistor_capacitor_sheet"]
        or record["in_finished_goods_sheet"]
        or record["in_historical_inventory"]
        for record in records
    )
    incremental_resistor_count = sum(
        not record["in_main_inventory_sheet"] and record["in_resistor_capacitor_sheet"] for record in records
    )
    overlap_main_resistor_count = sum(
        record["in_main_inventory_sheet"] and record["in_resistor_capacitor_sheet"] for record in records
    )
    return {
        "records": records,
        "missing_main_inventory": missing_main,
        "summary": {
            "bom_material_count": len(records),
            "material_master_match_count": sum(record["in_material_master"] for record in records),
            "main_inventory_match_count": in_main_count,
            "resistor_capacitor_match_count": in_resistor_count,
            "finished_goods_match_count": in_finished_count,
            "historical_inventory_match_count": in_historical_count,
            "all_inventory_source_union_count": all_source_count,
            "incremental_resistor_capacitor_match_count": incremental_resistor_count,
            "main_and_resistor_capacitor_overlap_count": overlap_main_resistor_count,
            "main_inventory_unmatched_count": len(missing_main),
        },
    }


def write_markdown(audit: dict[str, object], output_path: Path = DOC_OUTPUT) -> None:
    summary = audit["summary"]
    missing_main = audit["missing_main_inventory"]
    lines = [
        "# BOM 库存来源口径审计",
        "",
        "## 审计结论",
        "",
        "早期的 61/66 是以料号是否出现在主库存表或其他库存类工作表作为匹配口径；第五阶段的 44/66 则严格限定为一期允许使用的 `库存表`。两者不是同一个 MRP 可用库存口径。",
        "",
        "| 指标 | 结果 |",
        "| --- | ---: |",
        f"| BOM 料号数 | {summary['bom_material_count']} |",
        f"| 料号主数据匹配 | {summary['material_master_match_count']}/{summary['bom_material_count']} |",
        f"| 一期主库存表匹配 | {summary['main_inventory_match_count']}/{summary['bom_material_count']} |",
        f"| 电阻电容现库位表匹配 | {summary['resistor_capacitor_match_count']}/{summary['bom_material_count']} |",
        f"| 成品库存表匹配 | {summary['finished_goods_match_count']}/{summary['bom_material_count']} |",
        f"| 历史可处理库存匹配 | {summary['historical_inventory_match_count']}/{summary['bom_material_count']} |",
        f"| 四类来源并集匹配 | {summary['all_inventory_source_union_count']}/{summary['bom_material_count']} |",
        f"| 电阻电容表带来的新增匹配 | {summary['incremental_resistor_capacitor_match_count']} |",
        f"| 主库存表与电阻电容表的 BOM 料号交集 | {summary['main_and_resistor_capacitor_overlap_count']} |",
        f"| 主库存表未匹配 | {summary['main_inventory_unmatched_count']} |",
        "",
        "因此，17 个额外差异料号全部来自 `电阻电容现库位`；它们使“来源出现”口径从 44 增至 61。另有 5 个料号不在任一审计库存来源中。",
        "",
        "## 来源使用边界",
        "",
        "- 一期 MRP 只读取 `库存表` 的 `物料代码` 与 `库存`，按料号汇总；其余字段只保留明细，不改变该规则。",
        "- `电阻电容现库位` 的料号存在性参与本审计；`other_source_stock` 使用其 `实物数量`，不使用同表的 `库存数量`，避免将计划/账面字段与实物字段混为一个可用库存数。",
        "- `成品库存表` 以 `物料代码` 与 `库存` 审计，`2023年至2025年形成的存货-可以处理` 以 `物料代码` 与 `库存数量` 审计；两者均不进入一期 MRP。",
        "- 主库存与电阻电容来源已有 7 个 BOM 料号交集，证明来源不能在未做归属、时点和去重对账前直接相加。即使某料号仅出现于电阻电容表，也不代表它可安全作为一期可用库存。",
        "",
        "## 主库存表未匹配料号的逐项说明",
        "",
        "| 料号 | 名称 | 其他库存来源 | 其他来源实物/库存数量 | 可能重复来源 | 可纳入一期可用库存 | 说明 |",
        "| --- | --- | --- | ---: | --- | --- | --- |",
    ]
    for item in missing_main:
        stock = "" if item["other_source_stock"] is None else item["other_source_stock"]
        lines.append(
            "| {material_code} | {material_name} | {other_inventory_sources} | {stock} | {possible_duplicate_inventory_source} | {safe_for_phase_one_available_stock} | {reason} |".format(
                stock=stock,
                **item,
            )
        )
    lines.extend(
        [
            "",
            "## 对一期 MRP 的影响",
            "",
            "本审计不修改 services、tools 或 Repository 的一期计算边界，不改变任何 MRP 数值。17 个电阻电容来源记录和 5 个无来源记录均继续按现有一期规则处理：不自动并入 `available`，也不静默补零。",
        ]
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--json-output", type=Path, required=True)
    args = parser.parse_args()
    audit = build_audit()
    write_markdown(audit)
    args.json_output.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(audit["summary"], ensure_ascii=False))


if __name__ == "__main__":
    main()
